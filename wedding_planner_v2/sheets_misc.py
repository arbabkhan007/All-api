"""Wedding Planner V2 — Extras / Miscellaneous sheets (Novality Store)."""
from datetime import date
from openpyxl.styles import Alignment, Font

from core import (MONEY, MONEY2, DATEF, INT, sheet_shell, section, put_label,
                  put_formula, put_input, stat_pair, merged_input, make_table,
                  status_cf, protect, YESNO, DONE_OPTS, GREEN, AMBER, RED,
                  GUEST_SHEET, GUEST_FIRST, GUEST_LAST)

TAB = "7FA6A0"
AL_C = Alignment(horizontal="center", vertical="center")


def honeymoon(wb):
    ws = wb.create_sheet("Honeymoon Planner")
    sheet_shell(ws, TAB, "🏝️  HONEYMOON PLANNER",
                "Destination, budget, itinerary & packing list", "F")
    for col, w in zip("BCDEF", (26, 16, 16, 13, 26)):
        ws.column_dimensions[col].width = w

    section(ws, 5, "THE TRIP", "B", "F")
    fields = [("Destination", "Amalfi Coast, Italy"),
              ("Travel dates", "29 May – 9 Jun 2027"),
              ("Travelers", "The newlyweds")]
    for i, (k, v) in enumerate(fields):
        put_label(ws, f"B{6 + i}", k)
        put_input(ws, f"C{6 + i}", v)
        ws.merge_cells(f"C{6 + i}:F{6 + i}")

    section(ws, 10, "BUDGET", "B", "F")
    items = ["Flights", "Lodging", "Food & drink", "Activities", "Buffer / other"]
    sample = [(i, v) for i, v in zip(items, (2400, 2100, 900, 700, 500))]
    cols = [
        dict(key="item", title="Item", width=26),
        dict(key="est", title="Estimated", numfmt=MONEY, align=AL_C),
        dict(key="actual", title="Actual", numfmt=MONEY, align=AL_C),
        dict(key="bal", title="Over / under 🔒", kind="formula", numfmt=MONEY, align=AL_C,
             formula=lambda r, L:
                 f'=IF(AND(${L["est"]}{r}="",${L["actual"]}{r}=""),"",'
                 f'N(${L["actual"]}{r})-N(${L["est"]}{r}))'),
        dict(key="notes", title="Notes", width=26),
    ]
    # NOTE: total row lands on row 17 — the Dashboard reads 'Honeymoon Planner'!$C$17
    _, _, _, total_row = make_table(ws, 11, cols, sample, 0,
                                    totals=("TOTAL", {"est": "sum",
                                                      "actual": "sum",
                                                      "bal": "sum"}))
    assert total_row == 17, f"honeymoon total moved to {total_row}"

    section(ws, 19, "ITINERARY", "B", "F")
    iti = [
        ("Day 1–2", "Positano", "Beach club & sunset dinner", "Yes"),
        ("Day 3–5", "Amalfi & Ravello", "Path of the Gods hike", "No"),
        ("Day 6–8", "Capri", "Boat day around the Faraglioni", "No"),
    ]
    icols = [
        dict(key="day", title="Day", width=12),
        dict(key="place", title="Where", width=16),
        dict(key="plan", title="Plan", width=30),
        dict(key="booked", title="Booked?", dv=YESNO, align=AL_C),
        dict(key="notes", title="Notes", width=26),
    ]
    make_table(ws, 20, icols, iti, 9, add_filter=False, freeze=False)

    section(ws, 34, "PACKING LIST", "B", "F")
    packing = [
        ("Passports & IDs", "Documents"), ("Booking confirmations", "Documents"),
        ("Travel insurance printout", "Documents"), ("Sunscreen & after-sun", "Health"),
        ("Medications", "Health"), ("Swimwear", "Clothing"),
        ("Walking shoes", "Clothing"), ("Evening outfits", "Clothing"),
        ("Light jackets", "Clothing"), ("Adapters & chargers", "Tech"),
        ("Camera & GoPro", "Tech"), ("E-reader & downloads", "Tech"),
        ("Reusable water bottles", "Other"), ("Travel pillow", "Other"),
        ("Snacks for travel days", "Other"), ("Currency / notify bank", "Other"),
    ]
    pcols = [
        dict(key="item", title="Item", width=26),
        dict(key="cat", title="Category", width=16),
        dict(key="packed", title="Packed?", dv=DONE_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=30),
    ]
    first, last, L, _ = make_table(ws, 35, pcols, packing, 2, add_filter=False,
                                   freeze=False)
    status_cf(ws, f"D{first}:D{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def name_change(wb):
    ws = wb.create_sheet("Name Change Checklist")
    sheet_shell(ws, TAB, "📛  NAME CHANGE CHECKLIST",
                "Work through the paperwork one win at a time", "F")
    for col, w in zip("BCDEF", (30, 34, 10, 14, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Social Security card", "SS-5 form + certified marriage certificate"),
        ("Driver's license", "New SS card + current license"),
        ("Passport", "DS-5504 form + certificate + photo"),
        ("Bank accounts", "ID + marriage certificate, in branch"),
        ("Credit & debit cards", "Call each issuer"),
        ("Employer / payroll", "HR form + new SS card"),
        ("Voter registration", "Re-register with new name"),
        ("Health insurance", "Qualifying life event update"),
        ("Car title & insurance", "Certificate + updated license"),
        ("Utilities & phone", "Just call — easy win"),
        ("Doctor & dentist offices", "Call ahead of next visit"),
        ("Professional licenses", "Varies by licensing board"),
        ("Loyalty programs & subscriptions", "Gradual clean-up"),
    ]
    cols = [
        dict(key="item", title="Who to update", width=30),
        dict(key="docs", title="Documents needed", width=34),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
        dict(key="date", title="Date completed", numfmt=DATEF, align=AL_C),
        dict(key="notes", title="Notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"D{first}:D{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def favors(wb):
    ws = wb.create_sheet("Favors Tracker")
    sheet_shell(ws, TAB, "🎁  WEDDING FAVOR IDEAS & TRACKER",
                "Ideas, quantities, spend & whether you picked it", "G")
    for col, w in zip("BCDEFG", (24, 20, 11, 9, 12, 10, 20)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Mini succulents", "Greenhouse Co.", 3.5, 90),
        ("Local honey jars", "Blue Ridge Bees", 6, 90),
        ("Custom candles", "Wick & Wonder", 7, 90),
        ("Seed packets", "Bloom & Twine", 2, 90),
        ("Bottle openers", "Etsy — MetalWorks", 4, 90),
        ("Coaster sets", "Etsy — Grain & Co.", 5, 90),
        ("Charity donation cards", "—", 1, 90),
        ("Custom matches", "Etsy — LitCo.", 1.5, 90),
    ]
    cols = [
        dict(key="idea", title="Favor idea", width=24),
        dict(key="supplier", title="Supplier", width=20),
        dict(key="each", title="Cost each", numfmt=MONEY2, align=AL_C),
        dict(key="qty", title="Qty", width=9, align=AL_C),
        dict(key="total", title="Total 🔒", kind="formula", numfmt=MONEY2, align=AL_C,
             formula=lambda r, L:
                 f'=IF(OR(${L["each"]}{r}="",${L["qty"]}{r}=""),"",'
                 f'${L["each"]}{r}*${L["qty"]}{r})'),
        dict(key="chosen", title="Chosen?", dv=YESNO, align=AL_C),
        dict(key="notes", title="Notes", width=20),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4,
                                   totals=("TOTAL", {"total": "sum"}))
    status_cf(ws, f"G{first}:G{last}", [("Yes", GREEN)])
    protect(ws)


def thank_you_cards(wb):
    ws = wb.create_sheet("Thank You Cards")
    sheet_shell(ws, TAB, "💌  THANK YOU CARD TRACKER",
                "Guest names pull in from the Guest List — track gifts & notes sent", "G")
    for col, w in zip("BCDEFG", (22, 26, 14, 14, 12, 22)):
        ws.column_dimensions[col].width = w

    stat_pair(ws, 5, "B", "CARDS SENT", '=COUNTIF($F$9:$F$208,"Sent")', INT)
    stat_pair(ws, 5, "D", "GIFTS LOGGED",
              '=SUMPRODUCT(($C$9:$C$208<>"")*1)', INT)

    sample = [
        (None, "Stand mixer", date(2027, 6, 2), date(2027, 6, 6)),
        (None, "Crystal decanter", date(2027, 6, 3), None),
        (None, "Honeymoon fund", date(2027, 6, 10), None),
    ]
    cols = [
        dict(key="guest", title="Guest 🔒", kind="formula", width=22,
             formula=lambda r, L:
                 f"=IF('{GUEST_SHEET}'!$B{r}=\"\",\"\",'{GUEST_SHEET}'!$B{r})"),
        dict(key="gift", title="Gift received", width=26),
        dict(key="received", title="Date received", numfmt=DATEF, align=AL_C),
        dict(key="sent", title="Thank-you sent", numfmt=DATEF, align=AL_C),
        dict(key="status", title="Status 🔒", kind="formula", align=AL_C,
             formula=lambda r, L:
                 f'=IF(${L["sent"]}{r}<>"","Sent",'
                 f'IF(${L["gift"]}{r}<>"","To write",""))'),
        dict(key="notes", title="Notes", width=22),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample,
                                   GUEST_LAST - GUEST_FIRST + 1 - len(sample))
    status_cf(ws, f"F{first}:F{last}", [("Sent", GREEN), ("To write", AMBER)])
    protect(ws)


def registry(wb):
    ws = wb.create_sheet("Registry Checklist")
    sheet_shell(ws, TAB, "🛒  REGISTRY CHECKLIST",
                "Stores, items & purchased status", "H")
    for col, w in zip("BCDEFGH", (20, 24, 26, 11, 11, 16, 18)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Crate & Barrel", "crateandbarrel.com/registry", "Stand mixer", 380, None),
        ("Crate & Barrel", "crateandbarrel.com/registry", "Dinnerware set", 210, None),
        ("Crate & Barrel", "crateandbarrel.com/registry", "Bedding set", 160, None),
        ("Amazon", "amazon.com/wedding", "Crystal decanter", 70, None),
        ("Amazon", "amazon.com/wedding", "Robot vacuum", 400, None),
        ("Honeymoon fund", "our-wedding-website.com", "Experiences fund", None, None),
    ]
    cols = [
        dict(key="store", title="Store", width=20),
        dict(key="link", title="Registry link", width=24),
        dict(key="item", title="Item", width=26),
        dict(key="price", title="Price", numfmt=MONEY, align=AL_C),
        dict(key="bought", title="Purchased?", dv=YESNO, align=AL_C),
        dict(key="by", title="Purchased by", width=16),
        dict(key="notes", title="Notes", width=18),
    ]
    make_table(ws, 8, cols, sample, 10, totals=("TOTAL VALUE", {"price": "sum"}))
    protect(ws)


def vow_worksheet(wb):
    ws = wb.create_sheet("Vow Worksheet")
    sheet_shell(ws, TAB, "✍️  VOW WRITING WORKSHEET",
                "Answer honestly — then shape these into your vows", "F")
    for col, w in zip("BCDEF", (26, 22, 22, 22, 22)):
        ws.column_dimensions[col].width = w

    prompts = [
        ("1. How we met & my first impression",
         "It was a rainy Tuesday and the coffee shop had one table left…"),
        ("2. What I love most about you",
         "The way you talk to strangers' dogs like old friends…"),
        ("3. The exact moment I knew",
         None),
        ("4. What I promise you (big and small)",
         "I promise to always split the last slice…"),
        ("5. The funny bit — make them laugh",
         None),
        ("6. Our closing line — the vow they'll quote",
         None),
    ]
    r = 5
    for label, seed in prompts:
        put_label(ws, f"B{r}", label)
        r += 1
        merged_input(ws, f"B{r}:F{r + 2}", seed)
        r += 4
    tip = ws[f"B{r}"]
    tip.value = ("💡  Tip: vows read best at 1–2 minutes each. Read them out loud "
                 "and cut anything you stumble over.")
    tip.font = Font(size=9, italic=True, color="8A5A66")
    protect(ws)


def hashtag_social(wb):
    ws = wb.create_sheet("Hashtag & Social")
    sheet_shell(ws, TAB, "#️⃣  WEDDING HASHTAG & SOCIAL PLAN", "Get taggin'", "E")
    for col, w in zip("BCDE", (30, 30, 26, 24)):
        ws.column_dimensions[col].width = w

    section(ws, 5, "THE DETAILS", "B", "E")
    put_label(ws, "B6", "Auto hashtag 🔒")
    put_formula(ws, "C6", '=IF(Dashboard!$C$13="","",Dashboard!$C$13)')
    put_label(ws, "B7", "Partner 1 handle")
    put_input(ws, "C7", "@alex.morgan")
    put_label(ws, "B8", "Partner 2 handle")
    put_input(ws, "C8", "@jordan.lee")
    put_label(ws, "B9", "Backup / fun hashtags")
    put_input(ws, "C9", "#MorganLeeSayIDo  #FinallyFrickinMarried")
    ws.merge_cells("C9:E9")

    section(ws, 11, "HASHTAG IDEAS", "B", "E")
    ideas = [
        ("#MorganLeeSayIDo", "Instagram, TikTok, signage"),
        ("#TheMorgansGetHitched", "Photo booth prints"),
        ("#EatDrinkBeMarried", "Late-night snack photos"),
        ("#AlexAndJordanTakeAmalfi", "Honeymoon spam"),
    ]
    icols = [
        dict(key="tag", title="Hashtag idea", width=30),
        dict(key="where", title="Where to use", width=30),
        dict(key="notes", title="Notes", width=26),
    ]
    ws.column_dimensions["D"].width = 26
    make_table(ws, 12, icols, ideas, 4, add_filter=False, freeze=False)

    section(ws, 21, "SOCIAL PLAN", "B", "E")
    plan = [
        ("Instagram", "@alex.morgan", "Teaser post when save-the-dates mail"),
        ("TikTok", "@jordan.lee", "Planning-behind-the-scenes series"),
        ("Facebook group", "Morgan–Lee Wedding", "Guest logistics & carpooling"),
        ("Wedding website", "withjoy.com/morgan-lee", "RSVPs, registry, FAQs"),
    ]
    pcols = [
        dict(key="platform", title="Platform", width=30),
        dict(key="handle", title="Handle / link", width=30),
        dict(key="plan", title="Plan", width=26),
    ]
    make_table(ws, 22, pcols, plan, 3, add_filter=False, freeze=False)

    section(ws, 31, "SIGNAGE & REMINDERS", "B", "E")
    signs = [
        ("Hashtag sign for reception", "Handmade paper co."),
        ("Unplugged ceremony sign", "DIY calligraphy"),
        ("Welcome sign with QR to site", "Poster shop"),
    ]
    scols = [
        dict(key="sign", title="Sign", width=30),
        dict(key="by", title="Made by", width=30),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
    ]
    first, last, L, _ = make_table(ws, 32, scols, signs, 2, add_filter=False,
                                   freeze=False)
    status_cf(ws, f"D{first}:D{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def legal(wb):
    ws = wb.create_sheet("Legal & Officiant")
    sheet_shell(ws, TAB, "📜  LEGAL & OFFICIANT REQUIREMENTS",
                "Marriage license, deadlines & witnesses — suggested dates auto-set", "G")
    for col, w in zip("BCDEFG", (34, 16, 10, 14, 18, 24)):
        ws.column_dimensions[col].width = w

    rows = [
        ("Book officiant & confirm credentials", -180),
        ("Check county marriage license rules", -120),
        ("Apply for the marriage license (both of you)", -60),
        ("Confirm witness requirements (usually 2, 18+)", -60),
        ("Pick up the license within its valid window", -21),
        ("Officiant completes & files paperwork on the day", 0),
        ("Pack license, IDs & vows for the wedding day", -7),
        ("Return signed license to the county clerk", 7),
        ("Order certified copies for name change", 21),
        ("Start the name-change paperwork", 30),
    ]
    sample = [(r[0],) for r in rows]
    FORMULAS = {}
    for i, (_t, off) in enumerate(rows):
        rr = 9 + i
        if off == 0:
            FORMULAS[rr] = '=IF(Dashboard!$C$8="","",Dashboard!$C$8)'
        else:
            FORMULAS[rr] = (f'=IF(Dashboard!$C$8="","",'
                            f'Dashboard!$C$8{off:+d})')
    cols = [
        dict(key="req", title="Requirement", width=34),
        dict(key="by", title="Suggested by 🔒", kind="formula", numfmt=DATEF, align=AL_C,
             formula=lambda r, L: FORMULAS.get(r, '=IF(Dashboard!$C$8="","","")')),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
        dict(key="date", title="Date done", numfmt=DATEF, align=AL_C),
        dict(key="who", title="Who handles", width=18),
        dict(key="notes", title="Notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"D{first}:D{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    tip = ws[f"B{last + 2}"]
    tip.value = ("⚠️  Every county is different — always confirm rules, fees and "
                 "waiting periods with your local clerk.")
    tip.font = Font(size=9, italic=True, color="A93226")
    protect(ws)


def build(wb):
    honeymoon(wb)
    name_change(wb)
    favors(wb)
    thank_you_cards(wb)
    registry(wb)
    vow_worksheet(wb)
    hashtag_social(wb)
    legal(wb)
