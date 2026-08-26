"""Wedding Planner V2 — Music & Entertainment + Photography sheets (Novality Store)."""
from openpyxl.styles import Alignment

from core import (MONEY, sheet_shell, section, put_label, merged_input,
                  make_table, status_cf, protect,
                  YESNO, DONE_OPTS, GREEN, AMBER, RED)

TAB = "8E9AAF"
AL_C = Alignment(horizontal="center", vertical="center")


def ceremony_songs(wb):
    ws = wb.create_sheet("Ceremony Songs")
    sheet_shell(ws, TAB, "🎻  CEREMONY SONG LIST",
                "Processional, interlude & recessional", "F")
    for col, w in zip("BCDEF", (28, 34, 26, 16, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Prelude (guests arrive)", "Here Comes the Sun — instrumental", "The Beatles",
         "DJ", "8–10 songs ready"),
        ("Family seating", "Landslide — instrumental", "Fleetwood Mac", "DJ", None),
        ("Wedding party processional", "Can't Help Falling in Love — Kina Grannis",
         "Elvis Presley (cover)", "DJ", None),
        ("Bridal processional", "At Last — string quartet arrangement", "Etta James",
         "String trio", None),
        ("Interlude / unity ritual", "A Thousand Years — instrumental", "Christina Perri",
         "DJ", None),
        ("Recessional", "Signed, Sealed, Delivered", "Stevie Wonder", "DJ", None),
        ("Postlude", "Sunday Morning", "Maroon 5", "DJ", None),
    ]
    cols = [
        dict(key="moment", title="Moment", width=28),
        dict(key="song", title="Song", width=34),
        dict(key="artist", title="Artist", width=26),
        dict(key="performed", title="Performed by",
             dv=["DJ", "Band", "String trio", "Singer", "Recorded"], align=AL_C),
        dict(key="notes", title="Notes", width=24),
    ]
    make_table(ws, 8, cols, sample, 5)
    protect(ws)


def reception_playlist(wb):
    ws = wb.create_sheet("Reception Playlist")
    sheet_shell(ws, TAB, "🎧  RECEPTION PLAYLIST",
                "Must-play and do-not-play — hand this to your DJ", "E")
    for col, w in zip("BCDE", (32, 24, 24, 26)):
        ws.column_dimensions[col].width = w

    section(ws, 5, "MUST-PLAY 🙌", "B", "E")
    must = [
        ("September", "Earth, Wind & Fire", "Dance floor opener"),
        ("Mr. Brightside", "The Killers", "Everyone screams it"),
        ("Dancing Queen", "ABBA", "Grandma approves"),
        ("Shut Up and Dance", "Walk the Moon", None),
        ("Cha Cha Slide", "DJ Casper", "For the cousins"),
        ("Thinking Out Loud", "Ed Sheeran", "Slow-dance moment"),
        ("Uptown Funk", "Mark Ronson ft. Bruno Mars", None),
    ]
    mcols = [
        dict(key="song", title="Song", width=32),
        dict(key="artist", title="Artist", width=24),
        dict(key="moment", title="Moment / vibe", width=24),
        dict(key="notes", title="Notes", width=26),
    ]
    make_table(ws, 6, mcols, must, 8, add_filter=False, freeze=False)

    section(ws, 19, "DO-NOT-PLAY 🚫", "B", "E")
    dont = [
        ("Chicken Dance", "Werner Thomas", "Hard no"),
        ("Macarena", "Los del Río", "See above"),
        ("The Hokey Pokey", "Traditional", "We're not 5"),
    ]
    dcols = [
        dict(key="song", title="Song", width=32),
        dict(key="artist", title="Artist", width=24),
        dict(key="reason", title="Reason", width=24),
    ]
    make_table(ws, 20, dcols, dont, 5, add_filter=False, freeze=False)

    section(ws, 29, "DJ NOTES", "B", "E")
    merged_input(ws, "B30:E34",
                 "Keep energy up after dinner; dip volume during toasts; "
                 "check in before every set change.")
    protect(ws)


def special_dances(wb):
    ws = wb.create_sheet("Special Dances")
    sheet_shell(ws, TAB, "💃  SPECIAL DANCE SONGS",
                "First dance, father-daughter, mother-son & more", "F")
    for col, w in zip("BCDEF", (26, 32, 26, 20, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Grand entrance", "Beautiful Day — U2", None, "The couple"),
        ("First dance", "Better Together — Jack Johnson", None, "The couple"),
        ("Father–daughter dance", "My Girl — The Temptations", None, "Alex & dad"),
        ("Mother–son dance", "In My Life — The Beatles", None, "Jordan & mom"),
        ("Wedding party dance", "24K Magic — Bruno Mars", None, "Wedding party"),
        ("Bouquet toss song", "Single Ladies — Beyoncé", None, "Singles"),
        ("Garter toss song", "You Sexy Thing — Hot Chocolate", None, None),
        ("Last dance", "Landslide — Fleetwood Mac", None, "Everyone"),
    ]
    cols = [
        dict(key="dance", title="Dance", width=26),
        dict(key="song", title="Song", width=32),
        dict(key="artist", title="Artist", width=26),
        dict(key="who", title="Who dances", width=20),
        dict(key="notes", title="Notes", width=24),
    ]
    make_table(ws, 8, cols, sample, 4)
    protect(ws)


def entertainment(wb):
    ws = wb.create_sheet("Entertainment Ideas")
    sheet_shell(ws, TAB, "🎪  ENTERTAINMENT IDEAS",
                "Photo booth, live painter, fireworks — pick your favourites", "G")
    for col, w in zip("BCDEFG", (24, 22, 11, 10, 11, 22)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Photo booth & props", "Snap Booth Co.", 550),
        ("360° video booth", "SpinCam", 700),
        ("Live event painter", "Anna Brush Studio", 900),
        ("Caricature artist", "Doodles & Co.", 400),
        ("Cold sparkler fountains", "FX Nights", 600),
        ("Fireworks display", "SkyLine Pyro", 2500),
        ("Lawn games", "DIY", 80),
        ("Dance floor props", "DIY", 45),
        ("Acoustic cocktail set", "The Velvet Keys", 800),
        ("Magician (cocktail hour)", "The Amazing Ray", 500),
        ("Late-night karaoke hour", "Silver Sparrow DJ", 200),
    ]
    cols = [
        dict(key="idea", title="Idea", width=24),
        dict(key="provider", title="Provider", width=22),
        dict(key="cost", title="Est. cost", numfmt=MONEY, align=AL_C),
        dict(key="booked", title="Booked?", dv=YESNO, align=AL_C),
        dict(key="priority", title="Priority", dv=["High", "Medium", "Low"], align=AL_C),
        dict(key="notes", title="Notes", width=22),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"E{first}:E{last}", [("Yes", GREEN)])
    status_cf(ws, f"F{first}:F{last}", [("High", RED), ("Medium", AMBER)])
    protect(ws)


def shot_list(wb):
    ws = wb.create_sheet("Shot List")
    sheet_shell(ws, TAB, "📸  PHOTOGRAPHY SHOT LIST",
                "Must-have photos & candid moments — share with your photographer", "E")
    for col, w in zip("BCDE", (18, 42, 15, 24)):
        ws.column_dimensions[col].width = w

    shots = [
        ("Getting ready", "Dress / outfit hanging shot", "Must-have"),
        ("Getting ready", "Shoes, jewellery & perfume flat lay", "Must-have"),
        ("Getting ready", "Hair & makeup in progress", "Nice to have"),
        ("Getting ready", "First look with parents", "Must-have"),
        ("Details", "Invitation suite styled", "Must-have"),
        ("Details", "Rings close-up", "Must-have"),
        ("Details", "Bouquet & florals", "Must-have"),
        ("Details", "Venue signage & tablescape", "Must-have"),
        ("Details", "Vows & keepsake items", "Nice to have"),
        ("Ceremony", "Processional — both partners' reactions", "Must-have"),
        ("Ceremony", "Wide shot of vows from aisle end", "Must-have"),
        ("Ceremony", "Ring exchange close-up", "Must-have"),
        ("Ceremony", "First kiss", "Must-have"),
        ("Ceremony", "Recessional with guests cheering", "Must-have"),
        ("Couple", "Golden-hour couple portraits", "Must-have"),
        ("Couple", "Candid laughing walk", "Must-have"),
        ("Couple", "Vitamin-sea? no — mountain overlook shot", "Nice to have"),
        ("Family", "Full family formal (both sides)", "Must-have"),
        ("Family", "Grandparents candid hugs", "Must-have"),
        ("Reception", "Grand entrance reaction", "Must-have"),
        ("Reception", "First dance", "Must-have"),
        ("Reception", "Toasts — reactions of couple & speakers", "Must-have"),
        ("Reception", "Cake cutting", "Must-have"),
        ("Reception", "Dance floor chaos shot 😄", "Must-have"),
        ("Reception", "Tablescapes & room glow", "Nice to have"),
        ("Send-off", "Sparkler tunnel", "Must-have"),
        ("Send-off", "Getaway car exit", "Must-have"),
    ]
    sample = [(c, s, d) for c, s, d in shots]
    cols = [
        dict(key="cat", title="Category", width=18),
        dict(key="shot", title="Shot", width=42),
        dict(key="must", title="Priority", dv=["Must-have", "Nice to have"], align=AL_C),
        dict(key="notes", title="Who / notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 8)
    status_cf(ws, f"D{first}:D{last}", [("Must-have", RED)])
    protect(ws)


def family_photos(wb):
    ws = wb.create_sheet("Family Photos")
    sheet_shell(ws, TAB, "👨‍👩‍👧  FAMILY PHOTO GROUPINGS",
                "Formal portrait combinations — efficient line-ups save time", "F")
    for col, w in zip("BCDEF", (6, 40, 34, 16, 12)):
        ws.column_dimensions[col].width = w

    groups = [
        ("Couple + both sets of parents", "Morgans & Lees"),
        ("Couple + bride's immediate family", "Morgans"),
        ("Couple + groom's immediate family", "Lees"),
        ("Couple + grandparents", "All grandparents"),
        ("Couple + bridesmaids & groomsmen", "Full wedding party"),
        ("Couple + bridesmaids", None),
        ("Couple + groomsmen", None),
        ("Couple + flower girl & ring bearer", None),
        ("Extended family — bride's side", "Aunts, uncles, cousins"),
        ("Extended family — groom's side", "Aunts, uncles, cousins"),
        ("Chosen family / best friends", None),
        ("Everyone — all guests", "Grand group shot"),
    ]
    sample = [(i + 1, g, p) for i, (g, p) in enumerate(groups)]
    cols = [
        dict(key="n", title="#", width=6, align=AL_C),
        dict(key="group", title="Grouping", width=40),
        dict(key="people", title="People", width=34),
        dict(key="timing", title="Timing", dv=["After ceremony", "Cocktail hour",
                                               "Reception"], align=AL_C),
        dict(key="taken", title="Taken?", dv=DONE_OPTS, align=AL_C),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"F{first}:F{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def video_moments(wb):
    ws = wb.create_sheet("Video Moments")
    sheet_shell(ws, TAB, "🎥  VIDEO MUST-CAPTURE MOMENTS", "For your videographer", "E")
    for col, w in zip("BCDE", (36, 15, 20, 26)):
        ws.column_dimensions[col].width = w

    moments = [
        ("Vows (audio is everything — mic the groom)", "Must-have"),
        ("Both partners getting ready", "Must-have"),
        ("First look", "Must-have"),
        ("Processional reaction", "Must-have"),
        ("Ring exchange close-up", "Must-have"),
        ("First kiss & recessional", "Must-have"),
        ("Toasts with guest reactions", "Must-have"),
        ("First dance (wide + close)", "Must-have"),
        ("Parent dances", "Must-have"),
        ("Cake cutting", "Nice to have"),
        ("Guest messages / advice booth", "Nice to have"),
        ("Sparkler send-off & getaway car", "Must-have"),
        ("B-roll: venue, décor, details", "Must-have"),
        ("Late-night dance floor energy", "Nice to have"),
    ]
    sample = [(m, p) for m, p in moments]
    cols = [
        dict(key="moment", title="Moment", width=36),
        dict(key="priority", title="Priority", dv=["Must-have", "Nice to have"], align=AL_C),
        dict(key="who", title="Who captures", width=20),
        dict(key="notes", title="Notes", width=26),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"C{first}:C{last}", [("Must-have", RED)])
    protect(ws)


def build(wb):
    ceremony_songs(wb)
    reception_playlist(wb)
    special_dances(wb)
    entertainment(wb)
    shot_list(wb)
    family_photos(wb)
    video_moments(wb)
