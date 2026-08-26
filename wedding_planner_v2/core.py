"""
Wedding Planner Spreadsheet — Version 2
Author: Novality Store

Core styling, layout helpers and shared constants for the workbook generator.
All formula cells are LOCKED; every sheet is protected with PASSWORD.
"""
import os
from datetime import date

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

PASSWORD = "premium"
AUTHOR = "Novality Store"
VERSION = "Version 2.0"
PRODUCT = "Wedding Planner Spreadsheet"

# ---------------------------------------------------------------- palette
ROSE        = "B76E79"   # rose gold — banners
PLUM        = "6E4C5B"   # deep plum
SUB_ROSE    = "F6E7EA"   # subtitle strip
SECTION     = "8A5A66"   # section bars
HEADER      = "C9919C"   # table headers
INPUT_BG    = "FFFBEF"   # cream — editable cells
FORMULA_BG  = "F3EFEF"   # warm grey — locked formula cells
TOTAL_BG    = "EBD9DD"   # totals rows
BORDER_C    = "DCC3C9"

TITLE_FILL   = PatternFill("solid", start_color=ROSE)
SUB_FILL     = PatternFill("solid", start_color=SUB_ROSE)
SECTION_FILL = PatternFill("solid", start_color=SECTION)
HEADER_FILL  = PatternFill("solid", start_color=HEADER)
INPUT_FILL   = PatternFill("solid", start_color=INPUT_BG)
FORMULA_FILL = PatternFill("solid", start_color=FORMULA_BG)
TOTAL_FILL   = PatternFill("solid", start_color=TOTAL_BG)

TITLE_FONT   = Font(name="Georgia", size=15, bold=True, color="FFFFFF")
SUB_FONT     = Font(name="Georgia", size=9, italic=True, color="A0707A")
SECTION_FONT = Font(name="Georgia", size=10.5, bold=True, color="FFFFFF")
HEADER_FONT  = Font(size=10, bold=True, color="FFFFFF")
INPUT_FONT   = Font(size=10, color="2E2A2B")
FORMULA_FONT = Font(size=10, italic=True, color="6E5F64")
TOTAL_FONT   = Font(size=10, bold=True, color="5C4048")
LABEL_FONT   = Font(size=10, bold=True, color="5C4048")

THIN   = Side(style="thin", color=BORDER_C)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AL_C  = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
AL_R  = Alignment(horizontal="right",  vertical="center")

MONEY  = '"$"#,##0'
MONEY2 = '"$"#,##0.00'
PCT    = '0%'
DATEF  = 'd mmm yyyy'
INT    = '#,##0'
DAYS   = '#,##0" days"'

UNLOCKED = Protection(locked=False)

# ------------------------------------------------- dropdown option lists
YESNO       = ["Yes", "No"]
YESNO_NA    = ["Yes", "No", "N/A"]
RSVP_OPTS   = ["Pending", "Confirmed", "Declined", "No Reply"]
MEAL_OPTS   = ["Chicken", "Beef", "Fish", "Vegetarian", "Vegan",
               "Gluten-Free", "Kids Meal", "TBD"]
CHECK_OPTS  = ["To Do", "In Progress", "Done", "N/A"]
DONE_OPTS   = ["Done", "Not Yet"]
PAY_OPTS    = ["Unpaid", "Deposit Paid", "Partially Paid", "Paid"]
CONTRACT_OPTS = ["Inquiry", "Quote Received", "Negotiating", "Booked",
                 "Contract Signed", "Declined"]
SHIP_OPTS   = ["To Order", "Ordered", "Shipped", "Received", "Altered", "Delivered"]
RATING_OPTS = ["1", "2", "3", "4", "5"]

GREEN = ("D9EAD9", "2E7D46")
RED   = ("F9D7D5", "A93226")
AMBER = ("FCF3CF", "B7791F")
BLUE  = ("DCE9F5", "2D5F8A")

# ---------------------------------------------- key cross-sheet anchors
GUEST_SHEET = "Guest List"
GUEST_FIRST, GUEST_LAST = 9, 208          # data rows on Guest List
MBUD_SHEET = "Master Budget"
MBUD_FIRST, MBUD_LAST, MBUD_TOTAL = 9, 30, 31
PAY_SHEET = "Payment Tracker"
PAY_FIRST, PAY_LAST = 9, 158


def ref(sheet, cell):
    """Quoted cross-sheet reference."""
    return f"'{sheet}'!{cell}"


# ------------------------------------------------------------ utilities
def _fill_row(ws, row, c1, c2, fill, font=None, height=None):
    for ci in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=ci)
        cell.fill = fill
        cell.border = BORDER
        if font:
            cell.font = font
    if height:
        ws.row_dimensions[row].height = height


def sheet_shell(ws, tab_color, title, subtitle, last_col):
    """Standard 3-row banner: title / author strip / legend."""
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 1.5
    ci2 = ws[f"B1"].column  # 2
    lidx = ws[f"{last_col}1"].column
    _fill_row(ws, 1, ci2, lidx, TITLE_FILL, TITLE_FONT, height=30)
    ws.merge_cells(f"B1:{last_col}1")
    ws["B1"] = title
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    _fill_row(ws, 2, ci2, lidx, SUB_FILL, SUB_FONT, height=15)
    ws.merge_cells(f"B2:{last_col}2")
    ws["B2"] = subtitle
    ws["B2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    lg = ws["B3"]
    lg.value = ("✏️  Cream cells = your input     🔒  Grey cells = protected "
                "formulas     •     " + AUTHOR + "  " + VERSION)
    lg.font = Font(size=8, italic=True, color="B58E96")
    ws.row_dimensions[3].height = 12
    ws.row_dimensions[4].height = 6


def section(ws, row, text, c1="B", c2="I"):
    ws.merge_cells(f"{c1}{row}:{c2}{row}")
    for ci in range(ws[f"{c1}1"].column, ws[f"{c2}1"].column + 1):
        cell = ws.cell(row=row, column=ci)
        cell.fill = SECTION_FILL
        cell.border = BORDER
    cell = ws[f"{c1}{row}"]
    cell.value = text
    cell.font = SECTION_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 20


def put_input(ws, coord, value=None, numfmt=None, align=None):
    c = ws[coord]
    c.fill = INPUT_FILL
    c.border = BORDER
    c.font = INPUT_FONT
    c.protection = UNLOCKED
    c.alignment = align or AL_L
    if value is not None:
        c.value = value
    if numfmt:
        c.number_format = numfmt
    return c


def put_formula(ws, coord, formula, numfmt=None, align=None, bold=False):
    c = ws[coord]
    c.fill = FORMULA_FILL
    c.border = BORDER
    c.font = Font(size=10, italic=True, bold=bold, color="55454B") if bold else FORMULA_FONT
    c.alignment = align or AL_L
    c.protection = Protection(locked=True)   # formulas are ALWAYS locked
    c.value = formula
    if numfmt:
        c.number_format = numfmt
    return c


def put_label(ws, coord, text, bold=True, align=None):
    c = ws[coord]
    c.font = LABEL_FONT if bold else Font(size=10, color="5C4048")
    c.alignment = align or Alignment(horizontal="left", vertical="center")
    if text is not None:
        c.value = text
    return c


def merged_input(ws, rng, value=None, numfmt=None):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    from openpyxl.utils.cell import range_boundaries
    mn_col, mn_row, mx_col, mx_row = range_boundaries(rng)
    for r in range(mn_row, mx_row + 1):
        for c in range(mn_col, mx_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.protection = UNLOCKED
    top = ws[first]
    top.font = INPUT_FONT
    top.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if value is not None:
        top.value = value
    if numfmt:
        top.number_format = numfmt
    return top


def stat_pair(ws, row, label_col, text, formula, numfmt=None, value_col=None):
    """Small stat: bold label + adjacent locked formula value."""
    from openpyxl.utils import column_index_from_string
    vc = value_col or get_column_letter(column_index_from_string(label_col) + 1)
    put_label(ws, f"{label_col}{row}", text)
    c = put_formula(ws, f"{vc}{row}", formula, numfmt=numfmt,
                    align=Alignment(horizontal="center", vertical="center"))
    return f"{vc}{row}"


def make_table(ws, hdr_row, cols, sample_rows, blank_rows,
               totals=None, add_filter=True, freeze=True):
    """
    Generic table builder. Starts at column B.
    cols: list of dicts: key,title,width,kind('input'|'formula'),numfmt,
          dv(list),formula(fn(row,L)->str),align
    totals: (label_text, {col_key: 'sum' | fn(row,L)->str} | None)
    Returns (first_data_row, last_data_row, L, total_row_or_None)
    """
    L = {c["key"]: get_column_letter(i + 2) for i, c in enumerate(cols)}

    for i, c in enumerate(cols):
        cl = get_column_letter(i + 2)
        ws.column_dimensions[cl].width = c.get("width", 14)
        h = ws.cell(row=hdr_row, column=i + 2, value=c["title"])
        h.fill = HEADER_FILL
        h.font = HEADER_FONT
        h.border = BORDER
        h.alignment = AL_C
    ws.row_dimensions[hdr_row].height = 26

    n = len(sample_rows)
    first = hdr_row + 1
    last = hdr_row + n + blank_rows

    for r in range(first, last + 1):
        sample = sample_rows[r - first] if (r - first) < n else None
        for i, c in enumerate(cols):
            cl = get_column_letter(i + 2)
            coord = f"{cl}{r}"
            if c.get("kind") == "formula":
                put_formula(ws, coord, c["formula"](r, L),
                            numfmt=c.get("numfmt"), align=c.get("align"))
            else:
                v = sample[i] if (sample and i < len(sample)) else None
                put_input(ws, coord, v, numfmt=c.get("numfmt"),
                          align=c.get("align", AL_L))

    total_row = None
    if totals:
        label_text, spec = totals
        total_row = last + 1
        for i, c in enumerate(cols):
            cl = get_column_letter(i + 2)
            cell = ws[f"{cl}{total_row}"]
            cell.fill = TOTAL_FILL
            cell.border = BORDER
            cell.font = TOTAL_FONT
            cell.alignment = AL_L
            if i == 0 and label_text:
                cell.value = label_text
            elif spec and c["key"] in spec:
                s = spec[c["key"]]
                if s == "sum":
                    cell.value = f"=SUM({cl}{first}:{cl}{last})"
                else:
                    cell.value = s(total_row, L)
                if c.get("numfmt"):
                    cell.number_format = c["numfmt"]

    for i, c in enumerate(cols):
        if c.get("dv"):
            dv = DataValidation(type="list",
                                formula1='"' + ",".join(c["dv"]) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            cl = get_column_letter(i + 2)
            dv.add(f"{cl}{first}:{cl}{last}")

    end_col = get_column_letter(len(cols) + 1)
    if add_filter:
        ws.auto_filter.ref = f"B{hdr_row}:{end_col}{last}"
    if freeze:
        ws.freeze_panes = f"C{hdr_row + 1}"
    return first, last, L, total_row


def status_cf(ws, rng, mapping):
    """mapping: list of (text, (bg, fg))"""
    for text, (bg, fg) in mapping:
        ws.conditional_formatting.add(
            rng,
            CellIsRule(operator="equal", formula=[f'"{text}"'],
                       fill=PatternFill("solid", start_color=bg),
                       font=Font(color=fg, bold=True)))


def protect(ws):
    """Lock the sheet so formula cells need the password to change."""
    ws.protection.sheet = True
    ws.protection.password = PASSWORD
    ws.protection.selectLockedCells = False   # allowed
    ws.protection.selectUnlockedCells = False
    ws.protection.formatCells = False         # allowed (restyling)
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False                # allowed
    ws.protection.autoFilter = False          # allowed
    ws.protection.insertRows = True           # blocked
    ws.protection.insertColumns = True
    ws.protection.deleteRows = True
    ws.protection.deleteColumns = True
    ws.protection.insertHyperlinks = True
    ws.oddFooter.right.text = f"{AUTHOR} • {PRODUCT} • {VERSION}"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = "B76E79"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def finish(wb, out_name=None):
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.title = f"{PRODUCT} — {VERSION}"
    wb.properties.subject = "Wedding planning"
    wb.properties.category = "Wedding Planner"
    wb.properties.keywords = "wedding, planner, budget, checklist, Novality Store"
    wb.properties.description = (f"{PRODUCT} {VERSION} by {AUTHOR}. "
                                 "Formula cells are protected.")
    wb.calculation.fullCalcOnLoad = True
    out = out_name or os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "..", "Wedding_Planner_V2_Novality_Store.xlsx")
    wb.save(out)
    return out
