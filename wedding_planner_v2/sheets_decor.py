"""Wedding Planner V2 — Décor & Design sheets (Novality Store)."""
from openpyxl.styles import PatternFill, Alignment

from core import (MONEY, sheet_shell, section, put_label, put_formula,
                  put_input, merged_input, make_table, status_cf, protect,
                  YESNO_NA, DONE_OPTS, GREEN, AMBER, BORDER, UNLOCKED,
                  INPUT_FILL, Font)

TAB = "D4B483"
AL_C = Alignment(horizontal="center", vertical="center")


def colors_theme(wb):
    ws = wb.create_sheet("Colors & Theme")
    sheet_shell(ws, TAB, "🎨  COLOR PALETTE & THEME",
                "Swatches, keywords and materials that define the look", "I")
    for col, w in zip("BCDEFGHI", (16, 16, 16, 16, 16, 16, 16, 16)):
        ws.column_dimensions[col].width = w

    section(ws, 5, "THE PALETTE — recolour any swatch to suit you", "B", "I")
    swatches = [("F7D8DE", "Blush"), ("E3A6B3", "Dusty rose"),
                ("B76E79", "Rose gold"), ("9CAF88", "Sage"),
                ("D4B483", "Gold"), ("EAD8C3", "Champagne"),
                ("FDF8F0", "Ivory"), ("6E4C5B", "Plum")]
    uses = ["Bridesmaid dresses", "Linens", "Stationery foil", "Greenery",
            "Candle holders", "Cake details", "Tablecloths", "Accent signage"]
    for i, (hexc, nm) in enumerate(swatches):
        cl = chr(ord("B") + i)
        sc = ws[f"{cl}6"]
        sc.fill = PatternFill("solid", start_color=hexc)
        sc.border = BORDER
        sc.protection = UNLOCKED
        ws.row_dimensions[6].height = 30
        put_input(ws, f"{cl}7", nm, align=AL_C)
        put_input(ws, f"{cl}8", uses[i], align=AL_C)
    put_label(ws, "B9", "↑ colour name", bold=False)
    put_label(ws, "B10", "↑ where it's used", bold=False)

    section(ws, 12, "THEME KEYWORDS", "B", "I")
    cols = [dict(key="kw", title="Keyword", width=16),
            dict(key="use", title="Use it for", width=44)]
    ws.column_dimensions["C"].width = 44
    kw_sample = [("Garden party", "Florals, arch, stationery borders"),
                 ("Heirloom", "Rings, cake topper, tableware"),
                 ("Candlelit", "Reception lighting, aisle glow"),
                 ("Organic", "Loose arrangements, flowing ribbon")]
    make_table(ws, 13, cols, kw_sample, 4, add_filter=False, freeze=False)

    section(ws, 21, "TEXTURES & MATERIALS", "B", "I")
    tx_cols = [dict(key="mat", title="Material", width=16),
               dict(key="where", title="Where it appears", width=44)]
    tx_sample = [("Linen", "Tablecloths & napkins"),
                 ("Velvet", "Ribbon on bouquets"),
                 ("Brass", "Candlesticks & arch detail"),
                 ("Handmade paper", "Menus & place cards")]
    make_table(ws, 22, tx_cols, tx_sample, 4, add_filter=False, freeze=False)

    section(ws, 30, "DESIGN NOTES", "B", "I")
    merged_input(ws, "B31:I36",
                 "Keep ceremony décor low so it never blocks the mountain view…")
    protect(ws)


def floral_planner(wb):
    ws = wb.create_sheet("Floral Planner")
    sheet_shell(ws, TAB, "🌸  FLORAL ARRANGEMENT PLANNER",
                "Bouquets, ceremony flowers, centerpieces & costs", "I")
    for col, w in zip("BCDEFGHI",
                      (22, 28, 8, 11, 12, 11, 11, 20)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Bridal bouquet", "Garden roses, ranunculus, trailing ribbon", 1, 180),
        ("Bridesmaid bouquets", "Smaller version of bridal", 4, 65),
        ("Boutonnieres", "Spray rose & greenery", 8, 12),
        ("Corsages", "Grandmothers & mothers", 4, 18),
        ("Ceremony arch", "Organic asymmetric install", 1, 850),
        ("Aisle arrangements", "Mason jar posies on shepherd hooks", 10, 35),
        ("Centerpieces", "Low compote + candles", 14, 55),
        ("Cake flowers", "Fresh blooms for tiers", 1, 45),
        ("Flower girl petals", "Blush freeze-dried petals", 2, 15),
        ("Extra greenery", "Eucalyptus bunches for signage", 6, 12),
    ]
    cols = [
        dict(key="item", title="Arrangement", width=22),
        dict(key="flowers", title="Flowers / details", width=28),
        dict(key="qty", title="Qty", width=8, align=AL_C),
        dict(key="each", title="Est. each", numfmt=MONEY, align=AL_C),
        dict(key="total", title="Est. total 🔒", kind="formula", numfmt=MONEY, align=AL_C,
             formula=lambda r, L:
                 f'=IF(OR(${L["qty"]}{r}="",${L["each"]}{r}=""),"",'
                 f'${L["qty"]}{r}*${L["each"]}{r})'),
        dict(key="actual", title="Actual", numfmt=MONEY, align=AL_C),
        dict(key="var", title="Variance 🔒", kind="formula", numfmt=MONEY, align=AL_C,
             formula=lambda r, L:
                 f'=IF(${L["actual"]}{r}="","",${L["actual"]}{r}-N(${L["total"]}{r}))'),
        dict(key="notes", title="Notes", width=20),
    ]
    make_table(ws, 8, cols, sample, 6,
               totals=("TOTAL", {"total": "sum", "actual": "sum", "var": "sum"}))
    protect(ws)


def ceremony_decor(wb):
    ws = wb.create_sheet("Ceremony Décor")
    sheet_shell(ws, TAB, "⛪  CEREMONY DÉCOR CHECKLIST",
                "Arch, aisle, signage & programs", "G")
    for col, w in zip("BCDEFG", (24, 22, 12, 12, 10, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Arbor / arch", "Bloom & Twine", 850, None),
        ("Aisle markers", "Bloom & Twine", 350, None),
        ("Aisle runner", "Rentals", 60, None),
        ("Welcome sign", "Handmade paper co.", 120, None),
        ("Program table", None, 40, None),
        ("Guest book & pen", None, 55, None),
        ("Candles & lanterns", "Party Perfect Rentals", 90, None),
        ("Entrance florals", "Bloom & Twine", 150, None),
        ("Unity ceremony setup", None, 25, None),
    ]
    cols = [
        dict(key="item", title="Item", width=24),
        dict(key="owner", title="Vendor / owner", width=22),
        dict(key="est", title="Est. cost", numfmt=MONEY, align=AL_C),
        dict(key="actual", title="Actual cost", numfmt=MONEY, align=AL_C),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 6,
                                   totals=("TOTAL", {"est": "sum", "actual": "sum"}))
    status_cf(ws, f"F{first}:F{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def reception_decor(wb):
    ws = wb.create_sheet("Reception Décor")
    sheet_shell(ws, TAB, "🥂  RECEPTION DÉCOR CHECKLIST",
                "Centerpieces, lighting, draping & backdrop", "G")
    for col, w in zip("BCDEFG", (24, 22, 12, 12, 10, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Centerpieces", "Bloom & Twine", 770, None),
        ("Table numbers", "Handmade paper co.", 45, None),
        ("Place cards", "Handmade paper co.", 60, None),
        ("Sweetheart table", "DIY", 80, None),
        ("Cake table styling", "DIY", 40, None),
        ("Dessert bar", None, 120, None),
        ("String lights / festoons", "Party Perfect Rentals", 220, None),
        ("Draping & backdrop", "Party Perfect Rentals", 350, None),
        ("Photo booth backdrop", None, 95, None),
        ("Lounge furniture", "Rentals", 260, None),
        ("Dance floor wrap", None, 150, None),
        ("Bar signage", None, 35, None),
    ]
    cols = [
        dict(key="item", title="Item", width=24),
        dict(key="owner", title="Vendor / owner", width=22),
        dict(key="est", title="Est. cost", numfmt=MONEY, align=AL_C),
        dict(key="actual", title="Actual cost", numfmt=MONEY, align=AL_C),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 6,
                                   totals=("TOTAL", {"est": "sum", "actual": "sum"}))
    status_cf(ws, f"F{first}:F{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def stationery(wb):
    ws = wb.create_sheet("Stationery Checklist")
    sheet_shell(ws, TAB, "✉️  STATIONERY SUITE CHECKLIST",
                "Invitations, RSVPs, menus, place cards, programs & signage", "H")
    for col, w in zip("BCDEFGH", (24, 22, 8, 11, 13, 12, 20)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Save-the-dates", "Handmade paper co.", 80, 180, None),
        ("Invitations & RSVP cards", "Handmade paper co.", 80, 340, None),
        ("Detail / info cards", "Handmade paper co.", 80, 90, None),
        ("Menus", "Handmade paper co.", 15, 60, None),
        ("Place cards", "Handmade paper co.", 90, 75, None),
        ("Programs", "Handmade paper co.", 80, 85, None),
        ("Welcome sign print", "Poster shop", 1, 55, None),
        ("Bar & dessert signs", "DIY", 3, 25, None),
        ("Thank-you cards", "Handmade paper co.", 80, 95, None),
        ("Stamps & envelopes", "Post office", 1, 120, None),
    ]
    cols = [
        dict(key="item", title="Item", width=24),
        dict(key="vendor", title="Vendor", width=22),
        dict(key="qty", title="Qty", width=8, align=AL_C),
        dict(key="cost", title="Cost", numfmt=MONEY, align=AL_C),
        dict(key="ordered", title="Ordered date", align=AL_C, numfmt="d mmm yyyy"),
        dict(key="delivered", title="Delivered?", dv=YESNO_NA, align=AL_C),
        dict(key="notes", title="Notes", width=20),
    ]
    make_table(ws, 8, cols, sample, 6, totals=("TOTAL", {"cost": "sum"}))
    protect(ws)


def build(wb):
    colors_theme(wb)
    floral_planner(wb)
    ceremony_decor(wb)
    reception_decor(wb)
    stationery(wb)
