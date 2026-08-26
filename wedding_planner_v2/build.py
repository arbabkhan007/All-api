"""
Wedding Planner Spreadsheet — Version 2  (author: Novality Store)

Build entry point: assembles all 52 tabs and saves the workbook.
Every formula cell is locked; every sheet is protected with the template
password (see core.PASSWORD).

Run:  python3 -m wedding_planner_v2.build
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook            # noqa: E402

import sheets_planning                   # noqa: E402
import sheets_budget                     # noqa: E402
import sheets_guests                     # noqa: E402
import sheets_vendors                    # noqa: E402
import sheets_attire                     # noqa: E402
import sheets_decor                      # noqa: E402
import sheets_catering                   # noqa: E402
import sheets_timeline                   # noqa: E402
import sheets_media                      # noqa: E402
import sheets_misc                       # noqa: E402
from core import finish                  # noqa: E402


def main():
    wb = Workbook()
    wb.remove(wb.active)

    sheets_planning.build(wb)
    sheets_budget.build(wb)
    sheets_guests.build(wb)
    sheets_vendors.build(wb)
    sheets_attire.build(wb)
    sheets_decor.build(wb)
    sheets_catering.build(wb)
    sheets_timeline.build(wb)
    sheets_media.build(wb)
    sheets_misc.build(wb)

    wb.active = 0
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "..", "Wedding_Planner_V2_Novality_Store.xlsx")
    saved = finish(wb, out)
    print(f"✔ saved {saved}  ({len(wb.sheetnames)} sheets)")
    print("  sheets:", ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
