"""Wedding Planner V2 — Attire & Beauty sheets (Novality Store)."""
from datetime import date
from openpyxl.styles import Alignment

from core import (MONEY, MONEY2, DATEF, INT, sheet_shell, put_label,
                  put_formula, make_table, status_cf, protect,
                  YESNO, SHIP_OPTS, RATING_OPTS, DONE_OPTS, GREEN, RED, AMBER)

TAB = "C3899B"
AL_C = Alignment(horizontal="center", vertical="center")


def dress_log(wb):
    ws = wb.create_sheet("Dress Shopping Log")
    sheet_shell(ws, TAB, "👰  BRIDE'S DRESS SHOPPING LOG",
                "Stores, styles, prices & the alterations timeline", "J")
    for col, w in zip("BCDEFGHIJ",
                      (20, 24, 11, 14, 9, 10, 13, 15, 22)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Ivory & Oak Bridal", "Maison Lilée — 'Amélie'", 1800, date(2026, 9, 14),
         "5", "Yes!", "Yes", date(2027, 4, 10), "THE one 💍"),
        ("Something Blue Atelier", "Vera Vow — 'Iris'", 2200, date(2026, 10, 2),
         "4", "Maybe", "Yes", None, "Loved the neckline"),
        ("Classic Corner Bridal", "Sofia Sage — 'Willow'", 1500, date(2026, 10, 18),
         "3", "No", "No", None, "Not me"),
    ]
    cols = [
        dict(key="store", title="Store / boutique", width=20),
        dict(key="style", title="Designer & style", width=24),
        dict(key="price", title="Price", numfmt=MONEY, align=AL_C),
        dict(key="appt", title="Appointment date", numfmt=DATEF, align=AL_C),
        dict(key="rating", title="Rating", dv=RATING_OPTS, align=AL_C),
        dict(key="yes", title="Say yes?", dv=["Yes!", "Maybe", "No"], align=AL_C),
        dict(key="alt", title="Alterations?", dv=YESNO, align=AL_C),
        dict(key="deadline", title="Alterations deadline", numfmt=DATEF, align=AL_C),
        dict(key="notes", title="Notes", width=22),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 12)
    status_cf(ws, f"G{first}:G{last}", [("Yes!", GREEN), ("No", RED), ("Maybe", AMBER)])
    protect(ws)


def bridesmaid_dresses(wb):
    ws = wb.create_sheet("Bridesmaid Dresses")
    sheet_shell(ws, TAB, "💃  BRIDESMAID DRESS TRACKER",
                "Sizes, colors, orders & delivery status", "I")
    for col, w in zip("BCDEFGHI",
                      (18, 24, 8, 10, 13, 15, 13, 20)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Olivia Carter", "Chiffon — Dusty rose", "8", 165, date(2027, 1, 10),
         date(2027, 3, 15), "Ordered"),
        ("Mia Delgado", "Chiffon — Dusty rose", "6", 165, date(2027, 1, 10),
         date(2027, 3, 15), "Shipped"),
        ("Grace Brooks", "Chiffon — Dusty rose", "12", 165, date(2027, 1, 12),
         date(2027, 3, 20), "To Order"),
        ("Chloe Fontaine", "Chiffon — Dusty rose", "10", 165, None,
         None, "To Order"),
    ]
    cols = [
        dict(key="name", title="Bridesmaid", width=18),
        dict(key="style", title="Dress style & color", width=24),
        dict(key="size", title="Size", width=8, align=AL_C),
        dict(key="price", title="Price", numfmt=MONEY, align=AL_C),
        dict(key="ordered", title="Ordered date", numfmt=DATEF, align=AL_C),
        dict(key="delivery", title="Expected delivery", numfmt=DATEF, align=AL_C),
        dict(key="status", title="Status", dv=SHIP_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=20),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 10)
    status_cf(ws, f"H{first}:H{last}",
              [("Received", GREEN), ("Delivered", GREEN), ("Shipped", AMBER),
               ("To Order", RED)])
    protect(ws)


def groom_attire(wb):
    ws = wb.create_sheet("Groom & Groomsmen")
    sheet_shell(ws, TAB, "🤵  GROOM & GROOMSMEN ATTIRE",
                "Suit / tux details, rental vs purchase & fittings", "J")
    for col, w in zip("BCDEFGHIJK",
                      (18, 16, 22, 14, 13, 10, 13, 13, 13, 16)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Jordan Lee", "Groom", "Midnight blue tuxedo", "Midnight", "Purchase",
         950, date(2027, 3, 6), date(2027, 5, 15), "Ordered"),
        ("Chris Thompson", "Best Man", "Charcoal suit", "Charcoal", "Rental",
         180, date(2027, 3, 6), date(2027, 5, 18), "Ordered"),
        ("Mateo Reyes", "Groomsman", "Charcoal suit", "Charcoal", "Rental",
         180, date(2027, 3, 6), date(2027, 5, 18), "To Order"),
        ("Ethan Brooks", "Groomsman", "Charcoal suit", "Charcoal", "Rental",
         180, date(2027, 3, 8), date(2027, 5, 18), "To Order"),
    ]
    cols = [
        dict(key="name", title="Name", width=18),
        dict(key="role", title="Role",
             dv=["Groom", "Best Man", "Groomsman", "Father of the Bride",
                 "Father of the Groom", "Ring Bearer", "Usher"], align=AL_C),
        dict(key="style", title="Suit / tux style", width=22),
        dict(key="color", title="Color", width=14),
        dict(key="mode", title="Rental or purchase", dv=["Rental", "Purchase"], align=AL_C),
        dict(key="price", title="Price", numfmt=MONEY, align=AL_C),
        dict(key="fitting", title="Fitting date", numfmt=DATEF, align=AL_C),
        dict(key="pickup", title="Pickup date", numfmt=DATEF, align=AL_C),
        dict(key="status", title="Status", dv=SHIP_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=16),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 10)
    status_cf(ws, f"J{first}:J{last}",
              [("Received", GREEN), ("Delivered", GREEN), ("Ordered", AMBER),
               ("To Order", RED)])
    protect(ws)


def hmu_trials(wb):
    ws = wb.create_sheet("Hair & Makeup Trials")
    sheet_shell(ws, TAB, "💅  HAIR & MAKEUP TRIAL NOTES",
                "Artists, looks, products & trial dates", "I")
    for col, w in zip("BCDEFGHI",
                      (20, 18, 18, 13, 30, 10, 14, 22)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Bride", "Priya Nair — Glow Studio", "Soft glam + loose waves",
         date(2027, 2, 20), "Dewy base · brown liner · rose lip", 150, "Yes"),
        ("Mother of the Bride", "Glow Studio", "Natural glam",
         date(2027, 3, 5), "Soft bronze eye", 85, "Yes"),
    ]
    cols = [
        dict(key="person", title="Person",
             dv=["Bride", "Mother of the Bride", "Mother of the Groom",
                 "Bridesmaid", "Flower girl"], width=20),
        dict(key="artist", title="Artist / studio", width=18),
        dict(key="service", title="Service / look", width=18),
        dict(key="date", title="Trial date", numfmt=DATEF, align=AL_C),
        dict(key="look", title="Look & products used", width=30),
        dict(key="price", title="Trial price", numfmt=MONEY, align=AL_C),
        dict(key="booked", title="Booked for the day?", dv=YESNO, align=AL_C),
        dict(key="notes", title="Notes", width=22),
    ]
    make_table(ws, 8, cols, sample, 8)
    protect(ws)


def beauty_timeline(wb):
    ws = wb.create_sheet("Beauty Timeline")
    sheet_shell(ws, TAB, "✨  BEAUTY TIMELINE",
                "Skincare, whitening, tanning & nails — dates auto-set from your wedding date",
                "E")
    for col, w in zip("BCDE", (30, 18, 11, 30)):
        ws.column_dimensions[col].width = w

    tasks = [
        ("Daily skincare routine starts", -6, "m"),
        ("Begin teeth whitening", -3, "m"),
        ("Final facial (no later!)", -14, "d"),
        ("Hair colour & trim", -30, "d"),
        ("Brow shaping", -7, "d"),
        ("Waxing appointments", -5, "d"),
        ("Spray tan trial", -21, "d"),
        ("Final spray tan", -2, "d"),
        ("Mani & pedi", -2, "d"),
        ("Eyelash extensions", -1, "d"),
    ]
    sample = [(t,) for t, _o, _u in tasks]
    FORMULAS = {}
    for i, (_t, off, unit) in enumerate(tasks):
        r = 9 + i
        if unit == "m":
            FORMULAS[r] = f'=IF(Dashboard!$C$8="","",EDATE(Dashboard!$C$8,{off}))'
        else:
            FORMULAS[r] = f'=IF(Dashboard!$C$8="","",Dashboard!$C$8{off})'
    cols = [
        dict(key="task", title="Beauty task", width=30),
        dict(key="start", title="Start by 🔒", kind="formula", numfmt=DATEF, align=AL_C,
             formula=lambda r, L: FORMULAS.get(r, '=IF(Dashboard!$C$8="","","")')),
        dict(key="done", title="Done?", dv=DONE_OPTS, align=AL_C),
        dict(key="notes", title="Notes", width=30),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4)
    status_cf(ws, f"D{first}:D{last}", [("Done", GREEN), ("Not Yet", AMBER)])
    protect(ws)


def build(wb):
    dress_log(wb)
    bridesmaid_dresses(wb)
    groom_attire(wb)
    hmu_trials(wb)
    beauty_timeline(wb)
