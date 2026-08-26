"""Wedding Planner V2 — Catering & Menu sheets (Novality Store)."""
from datetime import date
from openpyxl.styles import Alignment

from core import (MONEY, MONEY2, DATEF, INT, sheet_shell, section, put_label,
                  put_formula, put_input, merged_input, make_table, status_cf,
                  protect, YESNO, RATING_OPTS, GREEN, AMBER, Font)

TAB = "C97B63"
AL_C = Alignment(horizontal="center", vertical="center")


def menu_planning(wb):
    ws = wb.create_sheet("Menu Planning")
    sheet_shell(ws, TAB, "🍽️  MENU PLANNING",
                "Appetizers, mains, sides & desserts", "G")
    for col, w in zip("BCDEFG", (13, 30, 16, 10, 18, 22)):
        ws.column_dimensions[col].width = w

    courses = ["Appetizer", "Appetizer", "Appetizer", "Main", "Main", "Main",
               "Side", "Side", "Dessert", "Dessert", "Kids meal", "Late-night"]
    sample = [(c,) for c in courses]
    cols = [
        dict(key="course", title="Course", width=13),
        dict(key="dish", title="Dish", width=30),
        dict(key="diet", title="Dietary tags",
             dv=["None", "Vegetarian", "Vegan", "Gluten-Free", "Nut-free",
                 "Dairy-free"], align=AL_C),
        dict(key="servings", title="Servings", numfmt=INT, align=AL_C),
        dict(key="source", title="Sourced from", width=18),
        dict(key="notes", title="Notes", width=22),
    ]
    make_table(ws, 8, cols, sample, 8)
    protect(ws)


def cake_design(wb):
    ws = wb.create_sheet("Cake Design")
    sheet_shell(ws, TAB, "🎂  CAKE DESIGN & FLAVOR NOTES", "Tiers, flavors & payment", "G")
    for col, w in zip("BCDEFG", (18, 26, 22, 20, 9, 22)):
        ws.column_dimensions[col].width = w

    section(ws, 5, "THE BAKERY", "B", "G")
    put_label(ws, "B6", "Bakery")
    put_input(ws, "C6", "Sugarplum Cakery")
    put_label(ws, "B7", "Delivery time")
    put_input(ws, "C7", "1:00 PM (before reception)")
    put_label(ws, "B8", "Delivery & setup fee")
    put_input(ws, "C8", 45, numfmt=MONEY, align=AL_C)

    section(ws, 10, "CAKE TIERS", "B", "G")
    sample = [
        ("Tier 1 (bottom)", "Lemon elderflower", "Vanilla bean cream",
         "Swiss meringue buttercream", 60),
        ("Tier 2", "Almond champagne", "Raspberry preserve",
         "Swiss meringue buttercream", 40),
        ("Tier 3 (top)", "Chocolate salted caramel", "Salted caramel",
         "Chocolate ganache", 20),
    ]
    cols = [
        dict(key="tier", title="Tier", width=16),
        dict(key="flavor", title="Flavor", width=26),
        dict(key="filling", title="Filling", width=22),
        dict(key="frosting", title="Frosting", width=20),
        dict(key="serves", title="Serves", numfmt=INT, align=AL_C),
        dict(key="notes", title="Notes", width=22),
    ]
    make_table(ws, 11, cols, sample, 1, add_filter=False, freeze=False)

    section(ws, 16, "PAYMENT", "B", "G")
    put_label(ws, "B17", "Price quoted")
    put_input(ws, "C17", 650, numfmt=MONEY, align=AL_C)
    put_label(ws, "B18", "Deposit paid")
    put_input(ws, "C18", 200, numfmt=MONEY, align=AL_C)
    put_label(ws, "B19", "Balance due 🔒")
    put_formula(ws, "C19", '=IF(AND($C$17="",$C$18=""),"",MAX(0,N($C$17)-N($C$18)))',
                numfmt=MONEY, align=AL_C)

    section(ws, 21, "DESIGN & SETUP NOTES", "B", "G")
    merged_input(ws, "B22:G26",
                 "Semi-naked finish with fresh blush blooms matching the bouquets; "
                 "gold leaf on tier 2; cake topper 'M & L'.")
    protect(ws)


def bar_drinks(wb):
    ws = wb.create_sheet("Bar & Drinks")
    sheet_shell(ws, TAB, "🍹  BAR & DRINK MENU PLANNER",
                "Signature cocktails, wine, beer & non-alcoholic", "H")
    for col, w in zip("BCDEFGH", (20, 26, 12, 11, 20, 16, 20)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Signature cocktail", "The Rosewood Fizz — gin, grapefruit, rosemary",
         120, 340, "Harvest & Vine Co.", "Cocktail hour"),
        ("Signature mocktail", "Blush Spritz — raspberry, elderflower, soda",
         60, 60, "Harvest & Vine Co.", "All night"),
        ("Beer", "Local IPA + lager kegs", 3, 420, "Blue Ridge Brewing", "Reception"),
        ("Wine — red", "Pinot noir (case of 12)", 2, 240, "Vine & Barrel", "Dinner"),
        ("Wine — white", "Sauvignon blanc (case of 12)", 2, 220, "Vine & Barrel", "Dinner"),
        ("Wine — sparkling", "Prosecco for toasts (case)", 1, 150, "Vine & Barrel", "Toasts"),
        ("Spirits & mixers", "Whiskey, vodka, tonic, citrus", 1, 380, "Harvest & Vine Co.",
         "Reception"),
        ("Soft drinks", "Sodas, lemonade, iced tea", 1, 90, "Harvest & Vine Co.", "All night"),
        ("Coffee & tea", "Decaf bar with late-night dessert", 1, 75, "Harvest & Vine Co.",
         "Dessert"),
    ]
    cols = [
        dict(key="type", title="Type", width=20),
        dict(key="item", title="Item / brand", width=26),
        dict(key="qty", title="Qty needed", width=12, align=AL_C),
        dict(key="cost", title="Est. cost", numfmt=MONEY, align=AL_C),
        dict(key="supplier", title="Supplier", width=20),
        dict(key="when", title="Served when", width=16),
        dict(key="notes", title="Notes", width=20),
    ]
    make_table(ws, 8, cols, sample, 5, totals=("TOTAL", {"cost": "sum"}))
    protect(ws)


def tasting_notes(wb):
    ws = wb.create_sheet("Tasting Notes")
    sheet_shell(ws, TAB, "🥄  TASTING APPOINTMENT NOTES",
                "Rate what you taste — then book the winner", "H")
    for col, w in zip("BCDEFGH", (14, 22, 32, 9, 13, 11, 24)):
        ws.column_dimensions[col].width = w

    sample = [
        (date(2027, 1, 22), "Harvest & Vine Co.",
         "Tasting menu: short rib, salmon, pear salad", "5", "5", "Yes",
         "Booked — short rib is unreal"),
        (date(2027, 2, 8), "Sugarplum Cakery",
         "Cake: lemon elderflower + chocolate caramel", "5", "4", "Yes",
         "Lemon tier won everyone over"),
    ]
    cols = [
        dict(key="date", title="Date", numfmt=DATEF, align=AL_C),
        dict(key="vendor", title="Vendor", width=22),
        dict(key="items", title="What we tasted", width=32),
        dict(key="taste", title="Taste", dv=RATING_OPTS, align=AL_C),
        dict(key="present", title="Presentation", dv=RATING_OPTS, align=AL_C),
        dict(key="winner", title="Winner?", dv=YESNO, align=AL_C),
        dict(key="notes", title="Notes", width=24),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 8)
    status_cf(ws, f"G{first}:G{last}", [("Yes", GREEN), ("No", AMBER)])
    protect(ws)


def late_night(wb):
    ws = wb.create_sheet("Late-Night Snacks")
    sheet_shell(ws, TAB, "🌙  LATE-NIGHT SNACK IDEAS",
                "The midnight save — ideas, quantities & costs", "G")
    for col, w in zip("BCDEFG", (24, 22, 10, 11, 10, 22)):
        ws.column_dimensions[col].width = w

    sample = [
        ("Slider truck", "Bun & Run", 75, 450),
        ("Pizza slices", "Slice of Heaven", 80, 320),
        ("Churros", "Sweet carts", 60, 210),
        ("Milkshakes", "Sugarplum Cakery", 50, 175),
        ("Ramen cups", "Self-serve station", 40, 60),
        ("Mini tacos", "Fork & Flame", 70, 280),
        ("Donut wall", "Glazed & Amused", 60, 240),
        ("Pretzel station", "Pop & Pretzel", 50, 150),
    ]
    cols = [
        dict(key="idea", title="Snack idea", width=24),
        dict(key="vendor", title="Vendor", width=22),
        dict(key="qty", title="Qty", width=10, align=AL_C),
        dict(key="cost", title="Est. cost", numfmt=MONEY, align=AL_C),
        dict(key="chosen", title="Chosen?", dv=YESNO, align=AL_C),
        dict(key="notes", title="Notes", width=22),
    ]
    first, last, L, _ = make_table(ws, 8, cols, sample, 4,
                                   totals=("TOTAL", {"cost": "sum"}))
    status_cf(ws, f"F{first}:F{last}", [("Yes", GREEN)])
    protect(ws)


def build(wb):
    menu_planning(wb)
    cake_design(wb)
    bar_drinks(wb)
    tasting_notes(wb)
    late_night(wb)
