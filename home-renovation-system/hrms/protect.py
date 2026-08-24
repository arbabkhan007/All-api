"""Lock formula cells and protect every sheet with password 'premium'.

Input / data-entry cells stay unlocked so buyers can run the system.
Formula, header, banner, KPI and dashboard cells cannot be typed over
without Review → Unprotect Sheet → premium.
"""

from __future__ import annotations

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection

SHEET_PASSWORD = "premium"
AUTHOR = "premium"

LOCKED = Protection(locked=True)
UNLOCKED = Protection(locked=False)

# Entirely calculated / branded surfaces — no casual typing
LOCK_ALL = {
    "Start Here",
    "Dashboard",
    "Finance",
    "Admin",
    "Shopping List",
    "Roles & Permissions",
}

# Extra empty rows unlocked for new records (input columns only)
INPUT_PAD_TO = 200


def _is_formula(val) -> bool:
    return isinstance(val, str) and val.startswith("=")


def apply_cell_locks(ws) -> dict:
    """Lock formulas/headers; unlock data-entry cells. Returns a small report."""
    name = ws.title
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    locked = unlocked = 0

    formula_cols = set()
    for row in ws.iter_rows(min_row=5, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if _is_formula(cell.value):
                formula_cols.add(cell.column)

    def set_lock(cell, lock: bool):
        nonlocal locked, unlocked
        if isinstance(cell, MergedCell):
            return
        cell.protection = LOCKED if lock else UNLOCKED
        if lock:
            locked += 1
        else:
            unlocked += 1

    # Rows 1–4 always locked (banner + headers)
    for r in range(1, 5):
        for c in range(1, max_col + 1):
            set_lock(ws.cell(r, c), True)

    if name in LOCK_ALL:
        for r in range(5, max_row + 1):
            for c in range(1, max_col + 1):
                set_lock(ws.cell(r, c), True)
        return {"sheet": name, "locked": locked, "unlocked": unlocked, "mode": "lock-all"}

    if name == "Settings":
        # Column A (setting name) + C (notes) locked; B (value) unlocked
        for r in range(5, max(max_row, 40) + 1):
            set_lock(ws.cell(r, 1), True)
            set_lock(ws.cell(r, 2), False)
            set_lock(ws.cell(r, 3), True)
        return {"sheet": name, "locked": locked, "unlocked": unlocked, "mode": "settings"}

    if name == "Lookups":
        for r in range(5, max(max_row, 80) + 1):
            for c in range(1, max_col + 1):
                set_lock(ws.cell(r, c), False)
        return {"sheet": name, "locked": locked, "unlocked": unlocked, "mode": "lookups"}

    # Registers: lock formula columns + any formula cell; unlock the rest
    pad_to = max(max_row, INPUT_PAD_TO)
    for r in range(5, pad_to + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            lock = c in formula_cols or _is_formula(cell.value)
            set_lock(cell, lock)

    return {
        "sheet": name,
        "locked": locked,
        "unlocked": unlocked,
        "mode": "register",
        "formula_cols": [get_column_letter(c) for c in sorted(formula_cols)],
    }


def protect_sheet(ws, password: str = SHEET_PASSWORD) -> None:
    ws.protection.sheet = True
    ws.protection.password = password
    ws.protection.enable()
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.insertRows = True
    ws.protection.insertHyperlinks = True
    ws.protection.deleteRows = False
    ws.protection.insertColumns = False
    ws.protection.deleteColumns = False
    ws.protection.formatCells = False
    ws.protection.formatColumns = True
    ws.protection.formatRows = True
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True
    ws.protection.objects = False
    ws.protection.scenarios = True
    ws.protection.pivotTables = True


def protect_workbook(wb, password: str = SHEET_PASSWORD) -> list:
    """Lock formulas, protect every sheet, stamp author = premium."""
    reports = []
    for ws in wb.worksheets:
        reports.append(apply_cell_locks(ws))
        protect_sheet(ws, password)

    # Do not lock workbook structure — buyers may still add a scratch sheet.
    wb.security = WorkbookProtection(lockStructure=False, lockWindows=False)
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.keywords = "premium, home renovation, spreadsheet, google sheets"
    wb.properties.category = "Spreadsheets"
    return reports
