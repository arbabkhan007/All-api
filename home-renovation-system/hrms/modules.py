"""Contractor ERP modules (Module Quick Links).

Excel- and Google-Sheets-safe: SUMIFS / COUNTIFS / INDEX / MATCH / IFERROR only.
No FILTER, no array MATCH, no SMALL(IF), no spilled range refs.
"""

from __future__ import annotations

from datetime import date, timedelta

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.utils import get_column_letter

from . import data as D
from . import extra_data as E
from .styles import (
    ACCENT,
    BG,
    CUR,
    CUR2,
    DANGER,
    DANGER_SOFT,
    DATE,
    DEC1,
    FORMULA_BG,
    INFO,
    INPUT_BG,
    INT,
    PCT,
    PCT1,
    PRIMARY,
    PRIMARY_MID,
    PRIMARY_SOFT,
    ROW_ALT,
    SECONDARY,
    SUCCESS,
    SUCCESS_SOFT,
    TEXT,
    TEXT_MUTED,
    THIN,
    WARNING,
    WARNING_SOFT,
    WHITE,
    align,
    fill,
    font,
    freeze,
    header_row,
    input_note,
    kpi_card,
    section_label,
    set_col_widths,
)


def _imp():
    """Late import so workbook.py can finish loading first."""
    from . import workbook as W

    return W


# ---------------------------------------------------------------------------
# Module Hub — matches the screenshot MODULE QUICK LINKS layout
# ---------------------------------------------------------------------------
HUB_ROWS = [
    ("Project Module", PRIMARY, [
        ("Project Master Data", "Projects"),
        ("Project Phases", "Project Phases"),
        ("Task Management", "Tasks"),
        ("Gantt Chart", "Gantt Chart"),
    ]),
    ("Client Module", PRIMARY_MID, [
        ("Client Master Data", "Client Master Data"),
        ("Client Communication", "Client Communication"),
        ("Client Satisfaction", "Client Satisfaction"),
    ]),
    ("Budget & Finance", SUCCESS, [
        ("Budget Planning", "Budget"),
        ("Cost Tracking", "Expenses"),
        ("Budget Dashboard", "Finance"),
        ("Change Orders", "Change Orders"),
    ]),
    ("Contractor & Vendor", PRIMARY_SOFT, [
        ("Contractor Master List", "Contractors"),
        ("Vendor Master Data", "Suppliers"),
        ("Contractor Performance", "Contractor Performance"),
        ("Contractor Payments", "Payments"),
    ]),
    ("Materials", WARNING, [
        ("Material Master List", "Materials"),
        ("Material Estimation", "Material Estimation"),
        ("Purchase Orders", "Purchase Orders"),
        ("Material Inventory", "Material Inventory"),
    ]),
    ("Labor", ACCENT, [
        ("Worker Master Data", "Worker Master Data"),
        ("Worker Attendance", "Worker Attendance"),
        ("Labor Cost Calculator", "Labor Cost Calc"),
        ("Productivity Tracker", "Productivity Tracker"),
    ]),
    ("Room Tracking", PRIMARY_MID, [
        ("Room Area Master", "Rooms"),
        ("Room Work Checklist", "Room Work Checklist"),
        ("Room Cost Summary", "Room Cost Summary"),
    ]),
    ("Design", ACCENT, [
        ("Design Requirements", "Design Studio"),
        ("Measurements Specs", "Measurements Specs"),
        ("Material Selection Board", "Material Selection"),
    ]),
    ("Equipment", INFO, [
        ("Equipment Inventory", "Equipment Inventory"),
        ("Equipment Usage Log", "Equipment Usage Log"),
        ("Equipment Maintenance", "Equipment Maintenance"),
    ]),
    ("Quality", SUCCESS, [
        ("Quality Standards", "Quality Standards"),
        ("Inspection Log", "Inspections"),
        ("Defect Snagging List", "Defect Snagging List"),
    ]),
    ("Safety", DANGER, [
        ("Safety Checklist", "Safety Checklist"),
        ("Incident Accident Log", "Incident Accident Log"),
        ("Permits Compliance", "Permits"),
    ]),
    ("Quotation & Invoice", WARNING, [
        ("Quotation Builder", "Quotes"),
        ("Invoice Management", "Invoice Management"),
        ("Payment Receipts", "Payment Receipts"),
    ]),
    ("Finance", SUCCESS, [
        ("Income Tracker", "Income Tracker"),
        ("Expense Tracker", "Expenses"),
        ("Profit Loss Project", "Profit Loss Project"),
        ("Cash Flow Tracker", "Cash Flow Tracker"),
    ]),
    ("Documents", PRIMARY, [
        ("Document Register", "Documents"),
        ("Contract Register", "Contract Register"),
    ]),
    ("Warranty", SUCCESS, [
        ("Warranty Register", "Warranties"),
        ("After-Sales Service", "After-Sales Service"),
    ]),
    ("Reports", PRIMARY_MID, [
        ("Project Status Report", "Project Status Report"),
        ("Weekly Progress Report", "Weekly Progress Report"),
        ("KPI Dashboard", "KPI Dashboard"),
    ]),
]


def build_module_hub(wb):
    W = _imp()
    ws = W._sheet(
        wb, "Module Hub", PRIMARY, 12,
        "Every module from the contractor operating system  ·  click a link to jump",
    )
    ws.merge_cells("A4:L4")
    ws["A4"] = "MODULE QUICK LINKS"
    ws["A4"].font = font(16, True, WHITE)
    ws["A4"].fill = fill(PRIMARY_SOFT)
    ws["A4"].alignment = align("left", "center")
    for c in range(1, 13):
        ws.cell(4, c).fill = fill(PRIMARY_SOFT)
        ws.cell(4, c).border = THIN
    ws.row_dimensions[4].height = 28

    r = 5
    for name, color, links in HUB_ROWS:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        cell = ws.cell(r, 1, name)
        cell.font = font(11, True, WHITE)
        cell.fill = fill(color)
        cell.alignment = align("left", "center")
        ws.cell(r, 2).fill = fill(color)
        ws.cell(r, 1).border = THIN
        ws.cell(r, 2).border = THIN
        col = 3
        for label, target in links:
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
            link = ws.cell(r, col, f"→  {label}")
            link.hyperlink = f"#'{target}'!A1"
            link.font = font(10, False, PRIMARY)
            link.fill = fill(WHITE)
            link.alignment = align("left", "center")
            ws.cell(r, col).border = THIN
            ws.cell(r, col + 1).fill = fill(WHITE)
            ws.cell(r, col + 1).border = THIN
            col += 2
        while col <= 12:
            ws.cell(r, col).fill = fill(BG)
            ws.cell(r, col).border = THIN
            col += 1
        ws.row_dimensions[r].height = 24
        r += 1

    r += 1
    section_label(ws, r, 1, "Also in this workbook", 12)
    extras = [
        (r + 1, 1, "Start Here", "Start Here"),
        (r + 1, 3, "Dashboard", "Dashboard"),
        (r + 1, 5, "AI Insights", "AI Insights"),
        (r + 1, 7, "Shopping List", "Shopping List"),
        (r + 1, 9, "Calendar", "Calendar"),
        (r + 1, 11, "Admin", "Admin"),
        (r + 2, 1, "Inventory", "Inventory"),
        (r + 2, 3, "Maintenance", "Maintenance"),
        (r + 2, 5, "Users", "Users"),
        (r + 2, 7, "Settings", "Settings"),
        (r + 2, 9, "Lookups", "Lookups"),
        (r + 2, 11, "Roles", "Roles & Permissions"),
    ]
    for rr, cc, label, target in extras:
        ws.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 1)
        cell = ws.cell(rr, cc, label)
        cell.hyperlink = f"#'{target}'!A1"
        cell.font = font(10, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.alignment = align("center", "center")
        ws.row_dimensions[rr].height = 20

    input_note(
        ws, r + 4, 1,
        "Ivory cells on registers are unlocked inputs. Sand-grey cells are locked formulas. "
        "Author: Novality store   ·   Unprotect password: premium   ·   Switch homes on Settings → Active Project ID.",
        12,
    )
    set_col_widths(ws, [16, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16])
    freeze(ws, "A5")
    return ws


# ---------------------------------------------------------------------------
# Registers
# ---------------------------------------------------------------------------
def build_project_phases(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Project Phases"]
    rows = [r[:] for r in E.PROJECT_PHASES]
    ws, last = W.build_list_sheet(
        wb, "Project Phases", PRIMARY,
        "Phase register  ·  duration / spent / remaining are formulas  ·  spent = paid expenses in the phase dates",
        headers, rows,
        [10, 12, 22, 12, 12, 12, 12, 12, 12, 10, 14, 12, 16, 22, 28],
        date_cols=[4, 5], money_cols=[7, 8, 9], pct_cols=[10],
        validations=[("=Phase_Status", "K5:K80")],
        status_col=11,
        status_map={
            "To Do": TEXT_MUTED, "Scheduled": WARNING, "In Progress": ACCENT,
            "Inspection": INFO, "Completed": SUCCESS, "Blocked": DANGER,
        },
        last_col=15,
    )
    for r in range(5, last + 1):
        ws.cell(r, 6).value = f'=IF($A{r}="","",$E{r}-$D{r}+1)'
        ws.cell(r, 6).number_format = INT
        ws.cell(r, 8).value = (
            f'=IF($A{r}="","",SUMIFS(Expenses!$H:$H,Expenses!$C:$C,$B{r},'
            f'Expenses!$J:$J,"Paid",Expenses!$B:$B,">="&$D{r},Expenses!$B:$B,"<="&$E{r}))'
        )
        ws.cell(r, 8).number_format = CUR
        ws.cell(r, 9).value = f'=IF($A{r}="","",$G{r}-$H{r})'
        ws.cell(r, 9).number_format = CUR
        for c in (6, 8, 9):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
        ws.cell(r, 7).fill = fill(INPUT_BG)
        ws.cell(r, 7).number_format = CUR
        ws.cell(r, 10).number_format = PCT
    return ws, last


def build_gantt_chart(wb):
    W = _imp()
    ws = W._sheet(
        wb, "Gantt Chart", ACCENT, 14,
        "Weekly Gantt of Project Phases  ·  gold header = this week  ·  D done · N now · P planned",
    )
    headers = ["Phase ID", "Project", "Phase", "Start", "End", "Status", "% Complete"]
    header_row(ws, 4, headers)
    n = len(E.PROJECT_PHASES)
    for i in range(n):
        r = 5 + i
        src = 5 + i
        ws.cell(r, 1).value = f"=IFERROR(INDEX('Project Phases'!A:A,{src}),\"\")"
        ws.cell(r, 2).value = f"=IFERROR(INDEX('Project Phases'!B:B,{src}),\"\")"
        ws.cell(r, 3).value = f"=IFERROR(INDEX('Project Phases'!C:C,{src}),\"\")"
        ws.cell(r, 4).value = f"=IFERROR(INDEX('Project Phases'!D:D,{src}),\"\")"
        ws.cell(r, 5).value = f"=IFERROR(INDEX('Project Phases'!E:E,{src}),\"\")"
        ws.cell(r, 6).value = f"=IFERROR(INDEX('Project Phases'!K:K,{src}),\"\")"
        ws.cell(r, 7).value = f"=IFERROR(INDEX('Project Phases'!J:J,{src}),\"\")"
        ws.cell(r, 4).number_format = DATE
        ws.cell(r, 5).number_format = DATE
        ws.cell(r, 7).number_format = PCT
        for c in range(1, 8):
            ws.cell(r, c).fill = fill(FORMULA_BG if i % 2 == 0 else "EBE4D6")
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
    last = 4 + n

    start = date(2026, 6, 1)
    weeks = 28
    g0 = 9
    ws.cell(3, g0, "WEEKLY GANTT  ·  gold = this week")
    ws.cell(3, g0).font = font(9, True, PRIMARY)
    ws.merge_cells(start_row=3, start_column=g0, end_row=3, end_column=g0 + weeks - 1)
    for w in range(weeks):
        col = g0 + w
        d0 = start + timedelta(days=7 * w)
        cell = ws.cell(4, col, d0)
        cell.number_format = "D-MMM"
        cell.font = font(8, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.alignment = align("center", "center")
        cell.border = THIN
        ws.column_dimensions[get_column_letter(col)].width = 5
        letter = get_column_letter(col)
        for r in range(5, last + 1):
            ws.cell(r, col).value = (
                f'=IF(OR($D{r}="",$E{r}=""),"",'
                f'IF(AND({letter}$4<=$E{r},{letter}$4+6>=$D{r}),'
                f'IF($F{r}="Completed","D",IF($F{r}="In Progress","N","P")),""))'
            )
            ws.cell(r, col).alignment = align("center", "center")
            ws.cell(r, col).font = font(8, True, WHITE)
            ws.cell(r, col).border = THIN
        ws.conditional_formatting.add(
            f"{letter}4",
            FormulaRule(
                formula=[f"AND({letter}4<=AsOfDate,{letter}4+6>=AsOfDate)"],
                fill=fill(ACCENT),
                font=font(8, True, WHITE),
            ),
        )
    g_range = f"I5:{get_column_letter(g0 + weeks - 1)}{last}"
    ws.conditional_formatting.add(g_range, FormulaRule(formula=['I5="D"'], fill=fill(SUCCESS), font=font(8, True, SUCCESS)))
    ws.conditional_formatting.add(g_range, FormulaRule(formula=['I5="N"'], fill=fill(ACCENT), font=font(8, True, ACCENT)))
    ws.conditional_formatting.add(g_range, FormulaRule(formula=['I5="P"'], fill=fill(SECONDARY), font=font(8, True, SECONDARY)))
    set_col_widths(ws, [10, 12, 22, 12, 12, 14, 12, 4])
    freeze(ws, "A5")
    return ws


def build_clients(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Client Master Data", PRIMARY_MID,
        "Homeowners and leads  ·  one row per client / property pairing",
        E.HEADERS_EXTRA["Client Master Data"], E.CLIENTS,
        [12, 24, 12, 32, 16, 24, 14, 12, 12, 12, 22, 14, 16, 12, 10, 40],
        date_cols=[14], money_cols=[12],
        validations=[("=Client_Type", "C5:C80"), ("=Client_Status", "J5:J80")],
        status_col=10,
        status_map={"Lead": INFO, "Active": SUCCESS, "On Hold": WARNING, "Complete": PRIMARY, "Inactive": TEXT_MUTED},
    )


def build_client_comms(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Client Communication", INFO,
        "Formal CRM log  ·  meetings, calls, emails, site visits  ·  chat lives on Messages",
        E.HEADERS_EXTRA["Client Communication"], E.CLIENT_COMMS,
        [10, 12, 12, 12, 12, 22, 52, 16, 12, 12, 12],
        date_cols=[2, 9],
        validations=[("=Comm_Type", "E5:E80")],
        status_col=10,
        status_map={"Open": WARNING, "Closed": SUCCESS},
    )


def build_satisfaction(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Client Satisfaction", SUCCESS,
        "Scores 1–5  ·  NPS −100 to 100  ·  walkthroughs and surveys",
        E.HEADERS_EXTRA["Client Satisfaction"], E.CLIENT_SAT,
        [12, 12, 12, 12, 14, 14, 10, 10, 14, 10, 10, 8, 44, 12],
        date_cols=[2],
        validations=[("=Sat_Channel", "F5:F80")],
        status_col=14,
        status_map={"Open": WARNING, "Closed": SUCCESS},
    )


def build_performance(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Contractor Performance"]
    rows = []
    for c in D.CONTRACTORS:
        rows.append([c[0], c[1], c[3]] + [None] * 9)
    ws, last = W.build_list_sheet(
        wb, "Contractor Performance", PRIMARY,
        "Live roll-up from Jobs + Contractors + Defect Snagging List",
        headers, rows,
        [14, 24, 16, 8, 12, 12, 12, 10, 10, 10, 12, 14],
        money_cols=[5, 6, 7], pct_cols=[8, 9, 10],
        last_col=12,
    )
    for r in range(5, last + 1):
        ws.cell(r, 4).value = f'=IF($A{r}="","",COUNTIF(Jobs!D:D,$A{r}))'
        ws.cell(r, 5).value = f'=IF($A{r}="","",SUMIF(Jobs!D:D,$A{r},Jobs!K:K))'
        ws.cell(r, 6).value = f'=IF($A{r}="","",SUMIF(Jobs!D:D,$A{r},Jobs!L:L))'
        ws.cell(r, 7).value = f'=IF($A{r}="","",$E{r}-$F{r})'
        ws.cell(r, 8).value = f'=IFERROR(INDEX(Contractors!I:I,MATCH($A{r},Contractors!A:A,0)),"")'
        ws.cell(r, 9).value = f'=IFERROR(INDEX(Contractors!J:J,MATCH($A{r},Contractors!A:A,0)),"")'
        ws.cell(r, 10).value = f'=IFERROR(INDEX(Contractors!M:M,MATCH($A{r},Contractors!A:A,0)),"")'
        ws.cell(r, 11).value = (
            f'=IF($A{r}="","",COUNTIFS(\'Defect Snagging List\'!H:H,'
            f'IFERROR(INDEX(Contractors!C:C,MATCH($A{r},Contractors!A:A,0)),""),'
            f'\'Defect Snagging List\'!J:J,"<>Closed"))'
        )
        ws.cell(r, 12).value = (
            f'=IF($J{r}="","",IF($J{r}>=0.93,"Preferred",IF($J{r}>=0.88,"Solid",IF($J{r}>=0.8,"Watch","Do Not Use"))))'
        )
        for c in range(4, 13):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c in (5, 6, 7):
                ws.cell(r, c).number_format = CUR
            if c in (8, 9, 10):
                ws.cell(r, c).number_format = PCT
    W._status_cf(ws, "L", 5, max(last, 80), {
        "Preferred": SUCCESS, "Solid": INFO, "Watch": WARNING, "Do Not Use": DANGER,
    })
    return ws, last


def build_estimates(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Material Estimation"]
    rows = [r[:] for r in E.ESTIMATES]
    ws, last = W.build_list_sheet(
        wb, "Material Estimation", WARNING,
        "Takeoff  ·  qty + waste, material $, labor $, total",
        headers, rows,
        [10, 12, 10, 12, 32, 8, 8, 10, 12, 12, 10, 12, 12, 12, 12, 22],
        money_cols=[10, 12, 13, 14, 15], pct_cols=[8],
        last_col=16,
    )
    for r in range(5, last + 1):
        ws.cell(r, 9).value = f'=IF($A{r}="","",$G{r}*(1+$H{r}))'
        ws.cell(r, 9).number_format = DEC1
        ws.cell(r, 13).value = f'=IF($A{r}="","",$I{r}*$J{r})'
        ws.cell(r, 13).number_format = CUR
        ws.cell(r, 14).value = f'=IF($A{r}="","",$K{r}*$L{r})'
        ws.cell(r, 14).number_format = CUR
        ws.cell(r, 15).value = f'=IF($A{r}="","",$M{r}+$N{r})'
        ws.cell(r, 15).number_format = CUR
        for c in (9, 13, 14, 15):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
        ws.cell(r, 8).number_format = PCT
    t = last + 2
    ws.cell(t, 12, "TOTAL").font = font(11, True, WHITE)
    ws.cell(t, 13, f"=SUM(M5:M{last})").number_format = CUR
    ws.cell(t, 14, f"=SUM(N5:N{last})").number_format = CUR
    ws.cell(t, 15, f"=SUM(O5:O{last})").number_format = CUR
    for c in range(12, 16):
        ws.cell(t, c).fill = fill(PRIMARY)
        ws.cell(t, c).font = font(11, True, WHITE)
        ws.cell(t, c).border = THIN
    return ws, last


def build_purchase_orders(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Purchase Orders"]
    rows = [r[:] for r in E.PURCHASE_ORDERS]
    ws, last = W.build_list_sheet(
        wb, "Purchase Orders", WARNING,
        "PO header  ·  total / balance formulas  ·  line items live on Materials (Receipt column)",
        headers, rows,
        [10, 12, 12, 24, 12, 12, 10, 10, 12, 12, 12, 12, 12, 22],
        date_cols=[2, 12, 13], money_cols=[6, 7, 8, 9, 10, 11],
        validations=[("=PO_Status", "E5:E80")],
        status_col=5,
        status_map={
            "Draft": TEXT_MUTED, "Sent": INFO, "Confirmed": WARNING,
            "Partial": ACCENT, "Received": SUCCESS, "Closed": PRIMARY, "Cancelled": DANGER,
        },
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 9).value = f'=IF($A{r}="","",$F{r}+$G{r}+$H{r})'
        ws.cell(r, 11).value = f'=IF($A{r}="","",$I{r}-$J{r})'
        ws.cell(r, 9).number_format = CUR
        ws.cell(r, 11).number_format = CUR
        ws.cell(r, 9).fill = fill(FORMULA_BG)
        ws.cell(r, 11).fill = fill(FORMULA_BG)
        ws.cell(r, 9).border = THIN
        ws.cell(r, 11).border = THIN
        ws.cell(r, 9).font = font(10)
        ws.cell(r, 11).font = font(10)
    return ws, last


def build_stock(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Material Inventory"]
    rows = [r[:] for r in E.STOCK]
    ws, last = W.build_list_sheet(
        wb, "Material Inventory", WARNING,
        "Yard / garage stock  ·  available = on hand − reserved  ·  value = available × unit cost",
        headers, rows,
        [10, 12, 28, 12, 16, 10, 10, 10, 12, 12, 12, 12, 12, 22],
        date_cols=[13], money_cols=[10, 11],
        validations=[("=Stock_Status", "L5:L80")],
        status_col=12,
        status_map={
            "In Stock": SUCCESS, "Low": WARNING, "Out": DANGER,
            "On Order": INFO, "Reserved": ACCENT,
        },
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 8).value = f'=IF($A{r}="","",$F{r}-$G{r})'
        ws.cell(r, 11).value = f'=IF($A{r}="","",$H{r}*$J{r})'
        ws.cell(r, 11).number_format = CUR
        ws.cell(r, 8).fill = fill(FORMULA_BG)
        ws.cell(r, 11).fill = fill(FORMULA_BG)
        ws.cell(r, 8).border = THIN
        ws.cell(r, 11).border = THIN
        ws.cell(r, 8).font = font(10)
        ws.cell(r, 11).font = font(10)
        ws.cell(r, 10).number_format = CUR2
    return ws, last


def build_workers(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Worker Master Data"]
    rows = [r[:] for r in E.WORKERS]
    ws, last = W.build_list_sheet(
        wb, "Worker Master Data", ACCENT,
        "Crew book  ·  hourly = daily / 8",
        headers, rows,
        [12, 16, 14, 12, 14, 16, 12, 12, 10, 12, 12, 18, 16],
        date_cols=[11], money_cols=[7, 8],
        validations=[("=Skill_Level", "D5:D80"), ("=Worker_Status", "J5:J80"), ("=Trade", "C5:C80")],
        status_col=10,
        status_map={"Active": SUCCESS, "On Leave": WARNING, "Terminated": DANGER},
        last_col=13,
    )
    for r in range(5, last + 1):
        ws.cell(r, 8).value = f'=IF($A{r}="","",$G{r}/8)'
        ws.cell(r, 8).number_format = CUR2
        ws.cell(r, 8).fill = fill(FORMULA_BG)
        ws.cell(r, 8).border = THIN
        ws.cell(r, 8).font = font(10)
        ws.cell(r, 7).number_format = CUR
    return ws, last


def build_attendance(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Worker Attendance"]
    rows = [r[:] for r in E.ATTENDANCE]
    ws, last = W.build_list_sheet(
        wb, "Worker Attendance", ACCENT,
        "Daily timesheet  ·  hourly and OT pulled from Worker Master Data",
        headers, rows,
        [10, 12, 12, 12, 12, 12, 10, 10, 10, 12, 12, 12, 22],
        date_cols=[2], money_cols=[8, 10, 11, 12],
        validations=[("=Attendance_Status", "E5:E200")],
        status_col=5,
        status_map={
            "Present": SUCCESS, "Absent": DANGER, "Half Day": WARNING,
            "Overtime": ACCENT, "Holiday": INFO, "Sick": TEXT_MUTED,
        },
        last_col=13,
    )
    for r in range(5, last + 1):
        ws.cell(r, 8).value = f'=IF($C{r}="","",IFERROR(INDEX(\'Worker Master Data\'!H:H,MATCH($C{r},\'Worker Master Data\'!A:A,0)),0))'
        ws.cell(r, 9).value = f'=IF($C{r}="","",IFERROR(INDEX(\'Worker Master Data\'!I:I,MATCH($C{r},\'Worker Master Data\'!A:A,0)),1.5))'
        ws.cell(r, 10).value = f'=IF($A{r}="","",$F{r}*$H{r})'
        ws.cell(r, 11).value = f'=IF($A{r}="","",$G{r}*$H{r}*$I{r})'
        ws.cell(r, 12).value = f'=IF($A{r}="","",$J{r}+$K{r})'
        for c in range(8, 13):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c != 9:
                ws.cell(r, c).number_format = CUR2
    return ws, last


def build_labor_cost(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Labor Cost Calc"]
    rows = [r[:] for r in E.LABOR_COST]
    ws, last = W.build_list_sheet(
        wb, "Labor Cost Calc", ACCENT,
        "Job-costed labor  ·  hourly from Worker Master  ·  burden on Settings-style 18%",
        headers, rows,
        [10, 12, 10, 12, 28, 12, 12, 10, 10, 10, 10, 12, 10, 12, 12],
        date_cols=[6], money_cols=[9, 12, 13, 14, 15], pct_cols=[11],
        last_col=15,
    )
    for r in range(5, last + 1):
        ws.cell(r, 9).value = f'=IF($D{r}="","",IFERROR(INDEX(\'Worker Master Data\'!H:H,MATCH($D{r},\'Worker Master Data\'!A:A,0)),0))'
        ws.cell(r, 10).value = f'=IF($D{r}="","",IFERROR(INDEX(\'Worker Master Data\'!I:I,MATCH($D{r},\'Worker Master Data\'!A:A,0)),1.5))'
        ws.cell(r, 12).value = f'=IF($A{r}="","",$G{r}*$I{r})'
        ws.cell(r, 13).value = f'=IF($A{r}="","",$H{r}*$I{r}*$J{r})'
        ws.cell(r, 14).value = f'=IF($A{r}="","",($L{r}+$M{r})*$K{r})'
        ws.cell(r, 15).value = f'=IF($A{r}="","",$L{r}+$M{r}+$N{r})'
        for c in (9, 10, 12, 13, 14, 15):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c != 10:
                ws.cell(r, c).number_format = CUR2
        ws.cell(r, 11).number_format = PCT
        ws.cell(r, 11).fill = fill(INPUT_BG)
    t = last + 2
    ws.cell(t, 11, "TOTAL").font = font(11, True, WHITE)
    ws.cell(t, 15, f"=SUM(O5:O{last})").number_format = CUR
    for c in (11, 15):
        ws.cell(t, c).fill = fill(PRIMARY)
        ws.cell(t, c).font = font(11, True, WHITE)
        ws.cell(t, c).border = THIN
    return ws, last


def build_productivity(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Productivity Tracker"]
    rows = [r[:] for r in E.PRODUCTIVITY]
    ws, last = W.build_list_sheet(
        wb, "Productivity Tracker", ACCENT,
        "Planned vs actual units and hours  ·  units/hr and variance are formulas",
        headers, rows,
        [10, 12, 12, 10, 16, 20, 12, 12, 8, 12, 12, 12, 12, 12],
        date_cols=[2], pct_cols=[13],
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 12).value = f'=IF(OR($A{r}="",$K{r}=0),"",$H{r}/$K{r})'
        ws.cell(r, 12).number_format = "0.00"
        ws.cell(r, 13).value = f'=IF(OR($A{r}="",$G{r}=0),"",($H{r}-$G{r})/$G{r})'
        ws.cell(r, 13).number_format = PCT1
        ws.cell(r, 12).fill = fill(FORMULA_BG)
        ws.cell(r, 13).fill = fill(FORMULA_BG)
        ws.cell(r, 12).border = THIN
        ws.cell(r, 13).border = THIN
        ws.cell(r, 12).font = font(10)
        ws.cell(r, 13).font = font(10)
    return ws, last


def build_room_checks(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Room Work Checklist", PRIMARY_MID,
        "Punch / closeout items per room  ·  filter by Room ID",
        E.HEADERS_EXTRA["Room Work Checklist"], E.ROOM_CHECKS,
        [10, 12, 10, 12, 36, 16, 14, 12, 12, 28, 22],
        date_cols=[8, 9],
        validations=[("=Checklist_Status", "G5:G200")],
        status_col=7,
        status_map={"Not Started": TEXT_MUTED, "In Progress": ACCENT, "Complete": SUCCESS, "N/A": INFO},
    )


def build_room_costs(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Room Cost Summary"]
    rows = []
    for rm in D.ROOMS:
        rows.append([rm[0], rm[1], rm[2]] + [None] * 11)
    ws, last = W.build_list_sheet(
        wb, "Room Cost Summary", PRIMARY,
        "Live room P&L  ·  area / budget / status from Rooms  ·  spend from Expenses",
        headers, rows,
        [10, 12, 20, 12, 12, 12, 12, 12, 12, 12, 12, 10, 10, 14],
        money_cols=[5, 6, 7, 8, 9, 10, 11, 12], pct_cols=[13],
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 4).value = f'=IFERROR(INDEX(Rooms!Q:Q,MATCH($A{r},Rooms!A:A,0)),0)'
        ws.cell(r, 5).value = f'=IFERROR(INDEX(Rooms!K:K,MATCH($A{r},Rooms!A:A,0)),0)'
        ws.cell(r, 6).value = f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!D:D,$A{r},Expenses!E:E,"Materials",Expenses!J:J,"Paid"))'
        ws.cell(r, 7).value = f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!D:D,$A{r},Expenses!E:E,"Labor",Expenses!J:J,"Paid"))'
        ws.cell(r, 8).value = f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!D:D,$A{r},Expenses!J:J,"Paid")-$F{r}-$G{r})'
        ws.cell(r, 9).value = f'=IF($A{r}="","",$F{r}+$G{r}+$H{r})'
        ws.cell(r, 10).value = f'=IFERROR(INDEX(Rooms!S:S,MATCH($A{r},Rooms!A:A,0)),0)'
        ws.cell(r, 11).value = f'=IF($A{r}="","",$E{r}-$I{r}-$J{r})'
        ws.cell(r, 12).value = f'=IF(OR($A{r}="",$D{r}=0),"",$I{r}/$D{r})'
        ws.cell(r, 13).value = f'=IF(OR($A{r}="",$E{r}=0),"",($I{r}+$J{r})/$E{r})'
        ws.cell(r, 14).value = f'=IFERROR(INDEX(Rooms!J:J,MATCH($A{r},Rooms!A:A,0)),"")'
        for c in range(4, 15):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c in (5, 6, 7, 8, 9, 10, 11, 12):
                ws.cell(r, c).number_format = CUR
            if c == 4:
                ws.cell(r, c).number_format = DEC1
            if c == 13:
                ws.cell(r, c).number_format = PCT1
    ws.conditional_formatting.add(
        f"M5:M{max(last, 80)}",
        ColorScaleRule(
            start_type="num", start_value=0, start_color=SUCCESS,
            mid_type="num", mid_value=0.9, mid_color=WARNING,
            end_type="num", end_value=1.2, end_color=DANGER,
        ),
    )
    return ws, last


def build_measurements(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Measurements Specs"]
    rows = [r[:] for r in E.MEASUREMENTS]
    ws, last = W.build_list_sheet(
        wb, "Measurements Specs", ACCENT,
        "Field measure  ·  area / perimeter / net area are formulas",
        headers, rows,
        [10, 12, 10, 18, 12, 12, 12, 12, 14, 12, 12, 22, 10, 10],
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 8).value = f'=IF($A{r}="","",IF($F{r}="",$E{r}*$G{r},$E{r}*$F{r}))'
        ws.cell(r, 9).value = f'=IF(OR($A{r}="",$F{r}=""),"",2*($E{r}+$F{r}))'
        ws.cell(r, 11).value = f'=IF($A{r}="","",$H{r}-$J{r})'
        for c in (8, 9, 11):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            ws.cell(r, c).number_format = DEC1
    return ws, last


def build_selection(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Material Selection", ACCENT,
        "Finish board  ·  considering → sampled → approved → installed",
        E.HEADERS_EXTRA["Material Selection"], E.SELECTIONS,
        [10, 12, 10, 12, 22, 14, 12, 14, 22, 12, 12, 12, 16, 24, 18],
        date_cols=[12], money_cols=[10],
        validations=[("=Sel_Status", "K5:K80")],
        status_col=11,
        status_map={
            "Considering": TEXT_MUTED, "Sampled": INFO, "Approved": SUCCESS,
            "Rejected": DANGER, "Installed": PRIMARY,
        },
    )


def build_equipment(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Equipment Inventory", INFO,
        "Owned and rented tools  ·  daily rate used by the usage log",
        E.HEADERS_EXTRA["Equipment Inventory"], E.EQUIPMENT,
        [10, 18, 12, 12, 16, 12, 12, 14, 12, 16, 18, 12, 12],
        date_cols=[12], money_cols=[7, 13],
        validations=[("=Equip_Status", "H5:H80"), ("=Equip_Condition", "I5:I80")],
        status_col=8,
        status_map={"Available": SUCCESS, "In Use": ACCENT, "Maintenance": WARNING, "Retired": TEXT_MUTED},
    )


def build_equip_usage(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Equipment Usage Log"]
    rows = [r[:] for r in E.EQUIP_USAGE]
    ws, last = W.build_list_sheet(
        wb, "Equipment Usage Log", INFO,
        "Hours on site  ·  cost = hours/8 × daily rate (owned rate may be 0)",
        headers, rows,
        [10, 12, 10, 12, 16, 10, 12, 12, 12, 12, 22],
        date_cols=[2], money_cols=[7, 8],
        last_col=11,
    )
    for r in range(5, last + 1):
        ws.cell(r, 7).value = f'=IF($C{r}="","",IFERROR(INDEX(\'Equipment Inventory\'!G:G,MATCH($C{r},\'Equipment Inventory\'!A:A,0)),0))'
        ws.cell(r, 8).value = f'=IF($A{r}="","",$F{r}/8*$G{r})'
        ws.cell(r, 7).number_format = CUR
        ws.cell(r, 8).number_format = CUR
        ws.cell(r, 7).fill = fill(FORMULA_BG)
        ws.cell(r, 8).fill = fill(FORMULA_BG)
        ws.cell(r, 7).border = THIN
        ws.cell(r, 8).border = THIN
        ws.cell(r, 7).font = font(10)
        ws.cell(r, 8).font = font(10)
    return ws, last


def build_equip_maint(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Equipment Maintenance"]
    rows = [r[:] for r in E.EQUIP_MAINT]
    ws, last = W.build_list_sheet(
        wb, "Equipment Maintenance", INFO,
        "Tool service  ·  days until due vs AsOfDate",
        headers, rows,
        [10, 10, 22, 12, 12, 10, 14, 12, 28, 12],
        date_cols=[4, 5], money_cols=[6],
        last_col=10,
    )
    for r in range(5, last + 1):
        ws.cell(r, 10).value = f'=IF(OR($A{r}="",$E{r}=""),"",$E{r}-AsOfDate)'
        ws.cell(r, 10).fill = fill(FORMULA_BG)
        ws.cell(r, 10).border = THIN
        ws.cell(r, 10).font = font(10)
    ws.conditional_formatting.add(
        f"J5:J{max(last, 80)}",
        FormulaRule(formula=["AND(J5<>\"\",J5<=14)"], fill=fill(WARNING_SOFT), font=font(10, True, WARNING)),
    )
    return ws, last


def build_quality_stds(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Quality Standards", SUCCESS,
        "Spec book  ·  what “done” means on this job",
        E.HEADERS_EXTRA["Quality Standards"], E.QUALITY_STDS,
        [10, 12, 40, 12, 22, 14, 16, 24],
    )


def build_snags(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Defect Snagging List"]
    rows = [r[:] for r in E.SNAGS]
    ws, last = W.build_list_sheet(
        wb, "Defect Snagging List", DANGER,
        "Punch / snag list  ·  Open Rank feeds the status report (Excel-safe INDEX/MATCH)",
        headers, rows,
        [10, 12, 12, 10, 16, 44, 12, 16, 12, 14, 12, 12, 24, 10],
        date_cols=[2, 9, 12], money_cols=[11],
        validations=[("=Defect_Severity", "G5:G200"), ("=Defect_Status", "J5:J200")],
        status_col=10,
        status_map={
            "Open": DANGER, "Assigned": WARNING, "In Progress": ACCENT,
            "Fixed": INFO, "Verified": SUCCESS, "Closed": PRIMARY,
        },
        last_col=14,
    )
    for r in range(5, last + 1):
        ws.cell(r, 14).value = (
            f'=IF(OR($A{r}="",$J{r}="Closed",$C{r}<>ActiveProject),"",'
            f'COUNTIFS($C$5:$C{r},ActiveProject,$J$5:$J{r},"<>Closed"))'
        )
        ws.cell(r, 14).fill = fill(FORMULA_BG)
        ws.cell(r, 14).border = THIN
        ws.cell(r, 14).font = font(10)
    W._status_cf(ws, "G", 5, max(last, 200), {
        "Critical": DANGER, "Major": ACCENT, "Minor": WARNING, "Cosmetic": INFO,
    })
    return ws, last


def build_safety(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Safety Checklist", DANGER,
        "Daily / weekly site safety  ·  fail requires a corrective action",
        E.HEADERS_EXTRA["Safety Checklist"], E.SAFETY,
        [10, 12, 12, 36, 12, 12, 20, 24, 16, 12, 12],
        date_cols=[2, 10],
        validations=[("=Safety_Result", "F5:F200")],
        status_col=6,
        status_map={"Pass": SUCCESS, "Fail": DANGER, "N/A": TEXT_MUTED, "Corrected": INFO},
    )


def build_incidents(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Incident Accident Log", DANGER,
        "Near miss through recordable  ·  keep even the small ones",
        E.HEADERS_EXTRA["Incident Accident Log"], E.INCIDENTS,
        [10, 12, 12, 16, 16, 44, 14, 22, 16, 12, 12, 18],
        date_cols=[2],
        validations=[("=Incident_Type", "D5:D80")],
        status_col=11,
        status_map={"Open": DANGER, "Closed": SUCCESS},
    )


def build_invoices(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Invoice Management"]
    rows = [r[:] for r in E.INVOICES]
    ws, last = W.build_list_sheet(
        wb, "Invoice Management", WARNING,
        "Client draws  ·  total / balance / aging are formulas",
        headers, rows,
        [10, 12, 12, 12, 40, 12, 10, 12, 12, 12, 12, 12, 10],
        date_cols=[2, 11], money_cols=[6, 7, 8, 9, 10],
        validations=[("=Invoice_Status", "L5:L80")],
        status_col=12,
        status_map={
            "Draft": TEXT_MUTED, "Sent": INFO, "Partial": WARNING,
            "Paid": SUCCESS, "Overdue": DANGER, "Void": TEXT_MUTED,
        },
        last_col=13,
    )
    for r in range(5, last + 1):
        ws.cell(r, 8).value = f'=IF($A{r}="","",$F{r}+$G{r})'
        ws.cell(r, 10).value = f'=IF($A{r}="","",$H{r}-$I{r})'
        ws.cell(r, 13).value = f'=IF($A{r}="","",IF($L{r}="Paid",0,IF($K{r}="","" ,AsOfDate-$K{r})))'
        for c in (8, 10, 13):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
        ws.cell(r, 8).number_format = CUR
        ws.cell(r, 10).number_format = CUR
    return ws, last


def build_receipts(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Payment Receipts", SUCCESS,
        "Money in  ·  tie each receipt to an invoice",
        E.HEADERS_EXTRA["Payment Receipts"], E.RECEIPTS,
        [10, 12, 12, 12, 12, 12, 16, 12, 22],
        date_cols=[2, 8], money_cols=[5],
        validations=[("=Payment_Method", "F5:F80")],
    )


def build_income(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Income Tracker", SUCCESS,
        "Client payments and deposits  ·  Planned / Paid  ·  rolls into P&L and cash flow",
        E.HEADERS_EXTRA["Income Tracker"], E.INCOME,
        [10, 12, 12, 16, 16, 28, 12, 10, 12, 12, 16],
        date_cols=[2], money_cols=[7],
        validations=[("=Income_Type", "D5:D80"), ("=Expense_Status", "I5:I80")],
        status_col=9,
        status_map={"Paid": SUCCESS, "Planned": INFO, "Committed": WARNING, "Void": TEXT_MUTED},
    )


def build_pnl(wb):
    W = _imp()
    headers = E.HEADERS_EXTRA["Profit Loss Project"]
    rows = [[p[0], p[2]] + [None] * 8 for p in D.PROJECTS]
    ws, last = W.build_list_sheet(
        wb, "Profit Loss Project", SUCCESS,
        "Income from Income Tracker  ·  costs from paid Expenses  ·  switch nothing — both projects listed",
        headers, rows,
        [12, 32, 14, 12, 12, 12, 12, 14, 12, 12],
        money_cols=[3, 4, 5, 6, 7, 8], pct_cols=[9],
        last_col=10,
    )
    for r in range(5, last + 1):
        ws.cell(r, 3).value = f'=IF($A{r}="","",SUMIFS(\'Income Tracker\'!G:G,\'Income Tracker\'!C:C,$A{r},\'Income Tracker\'!I:I,"Paid"))'
        ws.cell(r, 4).value = f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!C:C,$A{r},Expenses!E:E,"Labor",Expenses!J:J,"Paid"))'
        ws.cell(r, 5).value = f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!C:C,$A{r},Expenses!E:E,"Materials",Expenses!J:J,"Paid"))'
        ws.cell(r, 6).value = (
            f'=IF($A{r}="","",SUMIFS(Expenses!H:H,Expenses!C:C,$A{r},Expenses!J:J,"Paid")-$D{r}-$E{r})'
        )
        ws.cell(r, 7).value = f'=IF($A{r}="","",$D{r}+$E{r}+$F{r})'
        ws.cell(r, 8).value = f'=IF($A{r}="","",$C{r}-$G{r})'
        ws.cell(r, 9).value = f'=IF(OR($A{r}="",$C{r}=0),"",$H{r}/$C{r})'
        ws.cell(r, 10).value = (
            f'=IF($A{r}="","",IF($H{r}<0,"Loss",IF($I{r}<0.1,"Thin","Healthy")))'
        )
        for c in range(3, 11):
            ws.cell(r, c).fill = fill(FORMULA_BG)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c <= 8:
                ws.cell(r, c).number_format = CUR
            if c == 9:
                ws.cell(r, c).number_format = PCT1
    W._status_cf(ws, "J", 5, max(last, 20), {"Healthy": SUCCESS, "Thin": WARNING, "Loss": DANGER})
    return ws, last


def build_cashflow(wb):
    W = _imp()
    ws = W._sheet(
        wb, "Cash Flow Tracker", SUCCESS, 8,
        "Monthly cash  ·  opening rolls from prior closing  ·  Mar–Dec 2026",
    )
    header_row(ws, 4, E.HEADERS_EXTRA["Cash Flow Tracker"])
    months = [date(2026, m, 1) for m in range(3, 13)]
    for i, m in enumerate(months):
        r = 5 + i
        nxt = date(2026, m.month + 1, 1) if m.month < 12 else date(2027, 1, 1)
        ws.cell(r, 1, m.strftime("%b %Y"))
        if i == 0:
            ws.cell(r, 2, 0)
        else:
            ws.cell(r, 2, f"=F{r-1}")
        ws.cell(r, 3).value = (
            f'=SUMIFS(\'Income Tracker\'!G:G,\'Income Tracker\'!I:I,"Paid",'
            f'\'Income Tracker\'!B:B,">="&DATE({m.year},{m.month},1),'
            f'\'Income Tracker\'!B:B,"<"&DATE({nxt.year},{nxt.month},1))'
        )
        ws.cell(r, 4).value = (
            f'=SUMIFS(Expenses!H:H,Expenses!J:J,"Paid",'
            f'Expenses!B:B,">="&DATE({m.year},{m.month},1),'
            f'Expenses!B:B,"<"&DATE({nxt.year},{nxt.month},1))'
        )
        ws.cell(r, 5).value = f"=C{r}-D{r}"
        ws.cell(r, 6).value = f"=B{r}+E{r}"
        ws.cell(r, 7, "As of 24 Aug 2026" if m.month == 8 else "")
        for c in range(1, 8):
            ws.cell(r, c).fill = fill(WHITE if i % 2 == 0 else ROW_ALT)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)
            if c >= 2 and c <= 6:
                ws.cell(r, c).number_format = CUR
                ws.cell(r, c).fill = fill(FORMULA_BG)
        ws.cell(r, 2).fill = fill(FORMULA_BG if i else INPUT_BG)
    last = 4 + len(months)
    line = LineChart()
    line.title = "Monthly net cash"
    data = Reference(ws, min_col=5, min_row=4, max_row=last)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.width = 16
    line.height = 7
    line.y_axis.numFmt = CUR
    ws.add_chart(line, "A18")
    set_col_widths(ws, [14, 14, 14, 14, 14, 14, 28])
    freeze(ws, "A5")
    return ws


def build_contracts(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "Contract Register", PRIMARY,
        "GC, design, subcontract, supply, change orders  ·  signed flag + retention",
        E.HEADERS_EXTRA["Contract Register"], E.CONTRACTS,
        [12, 12, 14, 28, 12, 12, 12, 12, 10, 12, 28, 22],
        date_cols=[6, 7], money_cols=[5], pct_cols=[10],
        validations=[("=Contract_Type", "C5:C80"), ("=Contract_Status", "H5:H80")],
        status_col=8,
        status_map={
            "Draft": TEXT_MUTED, "Sent": WARNING, "Signed": INFO,
            "Active": SUCCESS, "Expired": DANGER, "Closed": PRIMARY,
        },
    )


def build_after_sales(wb):
    W = _imp()
    return W.build_list_sheet(
        wb, "After-Sales Service", SUCCESS,
        "Callbacks, warranty claims, handover consults  ·  the job after the job",
        E.HEADERS_EXTRA["After-Sales Service"], E.AFTER_SALES,
        [10, 12, 12, 12, 16, 36, 10, 16, 12, 14, 10, 22],
        date_cols=[2, 9], money_cols=[11],
        validations=[("=Service_Type", "E5:E80"), ("=Service_Status", "J5:J80"), ("=Priority", "G5:G80")],
        status_col=10,
        status_map={
            "Open": DANGER, "Scheduled": WARNING, "In Progress": ACCENT,
            "Closed": SUCCESS, "Escalated": ACCENT,
        },
    )


# ---------------------------------------------------------------------------
# Reports
# ---------------------------------------------------------------------------
def build_status_report(wb):
    W = _imp()
    ws = W._sheet(
        wb, "Project Status Report", PRIMARY, 12,
        "Printable status pack for the active project  ·  follows Settings → Active Project ID",
    )
    ws.merge_cells("A4:L4")
    ws["A4"] = '="PROJECT STATUS  ·  "&IFERROR(INDEX(Projects!C:C,MATCH(ActiveProject,Projects!A:A,0)),ActiveProject)&"  ·  "&TEXT(AsOfDate,"DD MMMM YYYY")'
    ws["A4"].font = font(16, True, PRIMARY)
    ws["A4"].fill = fill(BG)

    kpi_card(ws, 6, 1, "PHASE",
             '=IFERROR(INDEX(Projects!F:F,MATCH(ActiveProject,Projects!A:A,0)),"—")',
             '=IFERROR(INDEX(Projects!E:E,MATCH(ActiveProject,Projects!A:A,0)),"")', PRIMARY, 2, 3)
    kpi_card(ws, 6, 4, "PROGRESS",
             '=IFERROR(INDEX(Projects!P:P,MATCH(ActiveProject,Projects!A:A,0)),0)',
             "Percent complete", ACCENT, 2, 3)
    ws["D7"].number_format = PCT
    kpi_card(ws, 6, 7, "BUDGET / SPENT",
             '=IFERROR(INDEX(Projects!J:J,MATCH(ActiveProject,Projects!A:A,0)),0)',
             '=TEXT(SUMIFS(Expenses!H:H,Expenses!C:C,ActiveProject,Expenses!J:J,"Paid"),"$#,##0")&" spent"',
             PRIMARY, 2, 3)
    ws["G7"].number_format = CUR
    kpi_card(ws, 6, 10, "HEALTH",
             '=IFERROR(INDEX(Projects!U:U,MATCH(ActiveProject,Projects!A:A,0)),"—")',
             "On Track / Watch / Over", SUCCESS, 2, 3)

    section_label(ws, 10, 1, "OPEN SNAGS (active project)", 6)
    for i, h in enumerate(["Due", "Room", "Issue", "Severity", "Owner", "Status"], 1):
        cell = ws.cell(11, i, h)
        cell.font = font(9, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.border = THIN
    for i in range(6):
        r = 12 + i
        n = i + 1
        ws.cell(r, 1).value = f"=IFERROR(INDEX('Defect Snagging List'!I:I,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        ws.cell(r, 1).number_format = DATE
        ws.cell(r, 2).value = f"=IFERROR(INDEX('Defect Snagging List'!D:D,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        ws.cell(r, 3).value = f"=IFERROR(INDEX('Defect Snagging List'!F:F,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        ws.cell(r, 4).value = f"=IFERROR(INDEX('Defect Snagging List'!G:G,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        ws.cell(r, 5).value = f"=IFERROR(INDEX('Defect Snagging List'!H:H,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        ws.cell(r, 6).value = f"=IFERROR(INDEX('Defect Snagging List'!J:J,MATCH({n},'Defect Snagging List'!N:N,0)),\"\")"
        for c in range(1, 7):
            ws.cell(r, c).fill = fill(WHITE if i % 2 == 0 else ROW_ALT)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(9)

    section_label(ws, 10, 8, "PHASES", 5)
    for i, h in enumerate(["Phase", "Status", "%", "End"], 8):
        cell = ws.cell(11, i, h)
        cell.font = font(9, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.border = THIN
    for i, ph in enumerate(E.PROJECT_PHASES[:12]):
        r = 12 + i
        pid = ph[0]
        ws.cell(r, 8).value = f'=IFERROR(INDEX(\'Project Phases\'!C:C,MATCH("{pid}",\'Project Phases\'!A:A,0)),"")'
        ws.cell(r, 9).value = f'=IFERROR(INDEX(\'Project Phases\'!K:K,MATCH("{pid}",\'Project Phases\'!A:A,0)),"")'
        ws.cell(r, 10).value = f'=IFERROR(INDEX(\'Project Phases\'!J:J,MATCH("{pid}",\'Project Phases\'!A:A,0)),0)'
        ws.cell(r, 11).value = f'=IFERROR(INDEX(\'Project Phases\'!E:E,MATCH("{pid}",\'Project Phases\'!A:A,0)),"")'
        ws.cell(r, 10).number_format = PCT
        ws.cell(r, 11).number_format = DATE
        for c in range(8, 12):
            ws.cell(r, c).fill = fill(WHITE if i % 2 == 0 else ROW_ALT)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(9)

    section_label(ws, 25, 1, "COUNTS", 6)
    stats = [
        (26, "Open tasks", '=COUNTIFS(Tasks!B:B,ActiveProject,Tasks!H:H,"<>Completed")'),
        (27, "Overdue tasks", '=COUNTIFS(Tasks!B:B,ActiveProject,Tasks!Q:Q,"Overdue")'),
        (28, "Open snags", '=COUNTIFS(\'Defect Snagging List\'!C:C,ActiveProject,\'Defect Snagging List\'!J:J,"<>Closed")'),
        (29, "Safety fails (open)", '=COUNTIFS(\'Safety Checklist\'!C:C,ActiveProject,\'Safety Checklist\'!F:F,"Fail")'),
        (30, "POs open", '=COUNTIFS(\'Purchase Orders\'!C:C,ActiveProject,\'Purchase Orders\'!E:E,"<>Received",\'Purchase Orders\'!E:E,"<>Closed")'),
        (31, "Invoices unpaid $", '=SUMIFS(\'Invoice Management\'!J:J,\'Invoice Management\'!D:D,ActiveProject)'),
    ]
    for r, label, fml in stats:
        ws.cell(r, 1, label).font = font(10, True, PRIMARY)
        ws.cell(r, 2, fml).font = font(10)
        ws.cell(r, 1).fill = fill(WHITE)
        ws.cell(r, 2).fill = fill(FORMULA_BG)
        ws.cell(r, 1).border = THIN
        ws.cell(r, 2).border = THIN
        if "unpaid" in label:
            ws.cell(r, 2).number_format = CUR

    set_col_widths(ws, [18, 12, 36, 12, 16, 12, 14, 22, 14, 10, 12, 12])
    freeze(ws, "A4")
    return ws


def build_weekly_report(wb):
    W = _imp()
    ws = W._sheet(
        wb, "Weekly Progress Report", PRIMARY_MID, 10,
        "Week-of view around AsOfDate  ·  reprint every Monday",
    )
    ws.merge_cells("A4:J4")
    ws["A4"] = '="WEEKLY PROGRESS  ·  week of "&TEXT(AsOfDate-WEEKDAY(AsOfDate,2)+1,"DD MMM")&"  ·  "&ActiveProject'
    ws["A4"].font = font(16, True, PRIMARY)
    ws["A4"].fill = fill(BG)

    kpi_card(ws, 6, 1, "TASKS CLOSED (all time)",
             '=COUNTIFS(Tasks!B:B,ActiveProject,Tasks!H:H,"Completed")',
             "Completed on the register", SUCCESS, 2, 3)
    kpi_card(ws, 6, 4, "OPEN / OVERDUE",
             '=COUNTIFS(Tasks!B:B,ActiveProject,Tasks!H:H,"<>Completed")',
             '=COUNTIFS(Tasks!B:B,ActiveProject,Tasks!Q:Q,"Overdue")&" overdue"',
             WARNING, 2, 3)
    kpi_card(ws, 6, 7, "SPENT THIS MONTH",
             '=SUMIFS(Expenses!H:H,Expenses!C:C,ActiveProject,Expenses!J:J,"Paid",Expenses!B:B,">="&DATE(YEAR(AsOfDate),MONTH(AsOfDate),1),Expenses!B:B,"<"&DATE(YEAR(AsOfDate),MONTH(AsOfDate)+1,1))',
             "Paid expenses", ACCENT, 2, 3)
    ws["G7"].number_format = CUR

    section_label(ws, 10, 1, "THIS WEEK ON THE CALENDAR", 6)
    for i, h in enumerate(["Date", "Time", "Event", "Owner", "Status"], 1):
        cell = ws.cell(11, i, h)
        cell.font = font(9, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.border = THIN
    for i in range(8):
        r = 12 + i
        n = i + 1
        ws.cell(r, 1).value = f"=IFERROR(INDEX(Calendar!A:A,MATCH({n},Calendar!H:H,0)),\"\")"
        ws.cell(r, 1).number_format = DATE
        ws.cell(r, 2).value = f"=IFERROR(INDEX(Calendar!B:B,MATCH({n},Calendar!H:H,0)),\"\")"
        ws.cell(r, 3).value = f"=IFERROR(INDEX(Calendar!C:C,MATCH({n},Calendar!H:H,0)),\"\")"
        ws.cell(r, 4).value = f"=IFERROR(INDEX(Calendar!F:F,MATCH({n},Calendar!H:H,0)),\"\")"
        ws.cell(r, 5).value = f"=IFERROR(INDEX(Calendar!G:G,MATCH({n},Calendar!H:H,0)),\"\")"
        for c in range(1, 6):
            ws.cell(r, c).fill = fill(WHITE if i % 2 == 0 else ROW_ALT)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(9)

    section_label(ws, 10, 7, "ATTENDANCE THIS WEEK (sample)", 4)
    ws["G11"] = "Present days"
    ws["H11"] = '=COUNTIF(\'Worker Attendance\'!E:E,"Present")+COUNTIF(\'Worker Attendance\'!E:E,"Overtime")'
    ws["G12"] = "Total entries"
    ws["H12"] = '=COUNTA(\'Worker Attendance\'!A5:A200)'
    ws["G13"] = "Labor $ (calc)"
    ws["H13"] = "=SUM('Labor Cost Calc'!O:O)"
    ws["H13"].number_format = CUR
    for r in range(11, 14):
        ws.cell(r, 7).font = font(10, True, PRIMARY)
        ws.cell(r, 7).fill = fill(WHITE)
        ws.cell(r, 8).fill = fill(FORMULA_BG)
        ws.cell(r, 7).border = THIN
        ws.cell(r, 8).border = THIN

    section_label(ws, 22, 1, "NARRATIVE (type here — unlocked on this report if you unprotect)", 8)
    ws.merge_cells("A23:J26")
    ws["A23"] = (
        "Maplewood week of 24 Aug 2026: cabinet boxes continue; floor finish Tuesday; "
        "bath punch (haze) Thursday; hardware still in transit so doors wait. "
        "Contingency is under 10% — water feature stays on hold. Cabin Japandi board out for a decision."
    )
    ws["A23"].alignment = align("left", "top", True)
    ws["A23"].fill = fill(WHITE)
    ws["A23"].font = font(11)

    set_col_widths(ws, [16, 12, 28, 16, 12, 12, 22, 16, 12, 12])
    freeze(ws, "A4")
    return ws


def build_kpi_dashboard(wb):
    W = _imp()
    ws = W._sheet(
        wb, "KPI Dashboard", PRIMARY, 12,
        "Portfolio construction KPIs  ·  all projects unless noted",
    )
    cards = [
        (5, 1, "ACTIVE PROJECTS", '=COUNTIF(Projects!E:E,"Active")', "Status = Active", PRIMARY),
        (5, 4, "PORTFOLIO BUDGET", "=SUM(Projects!J:J)", "All projects", ACCENT),
        (5, 7, "CASH IN (PAID)", '=SUMIF(\'Income Tracker\'!I:I,"Paid",\'Income Tracker\'!G:G)', "Income Tracker", SUCCESS),
        (5, 10, "CASH OUT (PAID)", '=SUMIF(Expenses!J:J,"Paid",Expenses!H:H)', "Expenses", WARNING),
        (9, 1, "OPEN SNAGS", '=COUNTIFS(\'Defect Snagging List\'!J:J,"<>Closed",\'Defect Snagging List\'!A:A,"<>")-1', "Not closed", DANGER),
        (9, 4, "SAFETY FAILS", '=COUNTIF(\'Safety Checklist\'!F:F,"Fail")', "Need correction", DANGER),
        (9, 7, "OVERDUE TASKS", '=COUNTIF(Tasks!Q:Q,"Overdue")', "All projects", WARNING),
        (9, 10, "AR BALANCE", '=SUM(\'Invoice Management\'!J:J)', "Unpaid invoices", ACCENT),
        (13, 1, "CREW PRESENT", '=COUNTIF(\'Worker Attendance\'!E:E,"Present")', "Attendance rows", SUCCESS),
        (13, 4, "EQUIP IN USE", '=COUNTIF(\'Equipment Inventory\'!H:H,"In Use")', "On a job", INFO),
        (13, 7, "OPEN POs", '=COUNTIFS(\'Purchase Orders\'!E:E,"<>Received",\'Purchase Orders\'!E:E,"<>Closed",\'Purchase Orders\'!A:A,"<>")-1', "Not received", WARNING),
        (13, 10, "AVG SATISFACTION", '=AVERAGE(\'Client Satisfaction\'!G:G)', "Overall 1–5", SUCCESS),
    ]
    for r, c, label, val, sub, color in cards:
        mid = kpi_card(ws, r, c, label, val, sub, color, 2, 3)
        if any(k in label for k in ("BUDGET", "CASH", "AR ")):
            mid.number_format = CUR
        if "SATISFACTION" in label:
            mid.number_format = "0.0"

    section_label(ws, 17, 1, "PROJECT MARGIN", 6)
    for i, h in enumerate(["Project", "Income", "Cost", "Profit", "Margin"], 1):
        cell = ws.cell(18, i, h)
        cell.font = font(9, True, WHITE)
        cell.fill = fill(PRIMARY)
        cell.border = THIN
    for i, p in enumerate(D.PROJECTS):
        r = 19 + i
        pid = p[0]
        ws.cell(r, 1).value = f'=IFERROR(INDEX(\'Profit Loss Project\'!B:B,MATCH("{pid}",\'Profit Loss Project\'!A:A,0)),"")'
        ws.cell(r, 2).value = f'=IFERROR(INDEX(\'Profit Loss Project\'!C:C,MATCH("{pid}",\'Profit Loss Project\'!A:A,0)),0)'
        ws.cell(r, 3).value = f'=IFERROR(INDEX(\'Profit Loss Project\'!G:G,MATCH("{pid}",\'Profit Loss Project\'!A:A,0)),0)'
        ws.cell(r, 4).value = f'=IFERROR(INDEX(\'Profit Loss Project\'!H:H,MATCH("{pid}",\'Profit Loss Project\'!A:A,0)),0)'
        ws.cell(r, 5).value = f'=IFERROR(INDEX(\'Profit Loss Project\'!I:I,MATCH("{pid}",\'Profit Loss Project\'!A:A,0)),0)'
        for c in range(2, 5):
            ws.cell(r, c).number_format = CUR
        ws.cell(r, 5).number_format = PCT1
        for c in range(1, 6):
            ws.cell(r, c).fill = fill(WHITE if i % 2 == 0 else ROW_ALT)
            ws.cell(r, c).border = THIN
            ws.cell(r, c).font = font(10)

    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Income vs cost"
    data = Reference(ws, min_col=2, min_row=18, max_col=3, max_row=18 + len(D.PROJECTS))
    cats = Reference(ws, min_col=1, min_row=19, max_row=18 + len(D.PROJECTS))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.width = 14
    bar.height = 7
    bar.y_axis.numFmt = CUR
    ws.add_chart(bar, "G18")

    set_col_widths(ws, [28, 14, 14, 14, 12, 12, 16, 14, 12, 14, 12, 12])
    freeze(ws, "A4")
    return ws


def build_all(wb):
    """Create every new ERP sheet. Call after the original registers exist."""
    build_module_hub(wb)
    build_project_phases(wb)
    build_gantt_chart(wb)
    build_clients(wb)
    build_client_comms(wb)
    build_satisfaction(wb)
    build_performance(wb)
    build_estimates(wb)
    build_purchase_orders(wb)
    build_stock(wb)
    build_workers(wb)
    build_attendance(wb)
    build_labor_cost(wb)
    build_productivity(wb)
    build_room_checks(wb)
    build_room_costs(wb)
    build_measurements(wb)
    build_selection(wb)
    build_equipment(wb)
    build_equip_usage(wb)
    build_equip_maint(wb)
    build_quality_stds(wb)
    build_snags(wb)
    build_safety(wb)
    build_incidents(wb)
    build_invoices(wb)
    build_receipts(wb)
    build_income(wb)
    build_pnl(wb)
    build_cashflow(wb)
    build_contracts(wb)
    build_after_sales(wb)
    build_status_report(wb)
    build_weekly_report(wb)
    build_kpi_dashboard(wb)


SHEET_ORDER = [
    "Start Here", "Module Hub", "Dashboard", "KPI Dashboard",
    "Projects", "Project Phases", "Tasks", "Gantt Chart", "Timeline",
    "Client Master Data", "Client Communication", "Client Satisfaction", "Messages",
    "Budget", "Expenses", "Finance", "Change Orders",
    "Income Tracker", "Profit Loss Project", "Cash Flow Tracker",
    "Contractors", "Suppliers", "Contractor Performance", "Payments",
    "Materials", "Material Estimation", "Purchase Orders", "Material Inventory", "Shopping List",
    "Worker Master Data", "Worker Attendance", "Labor Cost Calc", "Productivity Tracker", "Jobs",
    "Rooms", "Room Work Checklist", "Room Cost Summary",
    "Design Studio", "Measurements Specs", "Material Selection", "AI Insights",
    "Equipment Inventory", "Equipment Usage Log", "Equipment Maintenance",
    "Quality Standards", "Inspections", "Defect Snagging List",
    "Safety Checklist", "Incident Accident Log", "Permits",
    "Quotes", "Invoice Management", "Payment Receipts",
    "Documents", "Contract Register",
    "Warranties", "After-Sales Service",
    "Inventory", "Maintenance", "Calendar", "Notifications",
    "Project Status Report", "Weekly Progress Report", "Admin",
    "Users", "Properties", "Roles & Permissions",
    "Settings", "Lookups", "Audit Log",
]
