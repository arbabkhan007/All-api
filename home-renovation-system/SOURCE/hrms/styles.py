"""Brand styles for the Home Renovation Management System workbook.

Visual direction: premium interior-design studio, not generic construction software.
Formulas and layouts are Excel + Google Sheets compatible.
"""

from __future__ import annotations

from copy import copy

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    NamedStyle,
    PatternFill,
    Protection,
    Side,
)
from openpyxl.styles.numbers import FORMAT_CURRENCY_USD_SIMPLE, FORMAT_PERCENTAGE_00
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# ---------------------------------------------------------------------------
# Color system (prompt §13)
# ---------------------------------------------------------------------------
PRIMARY = "16352F"  # Deep Forest Green
PRIMARY_MID = "1F4A42"
PRIMARY_SOFT = "2D5C52"
SECONDARY = "E8DDCB"  # Warm Sand
SECONDARY_DEEP = "D4C6AE"
ACCENT = "C86B4A"  # Terracotta
ACCENT_SOFT = "E8B09A"
BG = "F8F6F1"  # Soft Ivory
BG_WARM = "F3EEE4"
WHITE = "FFFFFF"
TEXT = "202522"  # Charcoal
TEXT_MUTED = "5C6560"
SUCCESS = "719B7A"  # Sage
SUCCESS_SOFT = "DCE8DE"
WARNING = "D99A3D"  # Amber
WARNING_SOFT = "F6E6C6"
DANGER = "C75C5C"  # Muted Red
DANGER_SOFT = "F3D6D6"
INFO = "4A6B73"
INFO_SOFT = "D5E2E5"
CARD = "FFFFFF"
ROW_ALT = "F3EEE4"
INPUT_BG = "FFFCF7"
FORMULA_BG = "F0EBE1"
GANTT = "C86B4A"
GANTT_DONE = "719B7A"
GANTT_NOW = "16352F"

THIN = Border(
    left=Side(style="thin", color="E0D6C6"),
    right=Side(style="thin", color="E0D6C6"),
    top=Side(style="thin", color="E0D6C6"),
    bottom=Side(style="thin", color="E0D6C6"),
)
HAIR = Border(
    left=Side(style="hair", color="E8DDCB"),
    right=Side(style="hair", color="E8DDCB"),
    top=Side(style="hair", color="E8DDCB"),
    bottom=Side(style="hair", color="E8DDCB"),
)
NONE = Border()

CUR = '"$"#,##0'
CUR2 = '"$"#,##0.00'
PCT = "0%"
PCT1 = "0.0%"
DATE = "YYYY-MM-DD"
INT = "#,##0"
DEC1 = "0.0"


def fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def font(
    size=11,
    bold=False,
    color=TEXT,
    name="Calibri",
    italic=False,
) -> Font:
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def align(h="left", v="center", wrap=False) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def apply_cell(
    cell,
    *,
    value=None,
    fill_hex=None,
    font_obj=None,
    align_obj=None,
    border_obj=THIN,
    num_fmt=None,
    wrap=None,
):
    if value is not None:
        cell.value = value
    if fill_hex:
        cell.fill = fill(fill_hex)
    if font_obj:
        cell.font = font_obj
    if align_obj:
        cell.alignment = align_obj
    elif wrap is not None:
        cell.alignment = align("left", "center", wrap)
    if border_obj is not None:
        cell.border = border_obj
    if num_fmt:
        cell.number_format = num_fmt
    return cell


def paint_range(ws, start_row, start_col, end_row, end_col, hex_color):
    f = fill(hex_color)
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(r, c).fill = f


def set_col_widths(ws, widths):
    """widths: list of numbers, or dict {letter: width} / {index: width}."""
    if isinstance(widths, dict):
        for key, w in widths.items():
            if isinstance(key, int):
                ws.column_dimensions[get_column_letter(key)].width = w
            else:
                ws.column_dimensions[str(key)].width = w
    else:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def hide_grid(ws):
    ws.sheet_view.showGridLines = False


def page_setup(ws, title, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5, header=0.25, footer=0.25)
    ws.oddHeader.left.text = "Novality Store  ·  Home Renovation System"
    ws.oddHeader.right.text = title
    ws.oddFooter.left.text = "Confidential  ·  Digital twin of the home"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def banner(ws, title, subtitle, last_col=12, tab_color=PRIMARY):
    """Two-row branded banner used on every sheet."""
    ws.sheet_properties.tabColor = tab_color
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    c1 = ws.cell(1, 1, title)
    c1.font = font(18, True, WHITE)
    c1.fill = fill(PRIMARY)
    c1.alignment = align("left", "center")
    c1.border = NONE
    c2 = ws.cell(2, 1, subtitle)
    c2.font = font(10, False, PRIMARY)
    c2.fill = fill(SECONDARY)
    c2.alignment = align("left", "center")
    c2.border = NONE
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = fill(PRIMARY)
        ws.cell(1, col).border = NONE
        ws.cell(2, col).fill = fill(SECONDARY)
        ws.cell(2, col).border = NONE
    ws.row_dimensions[3].height = 8
    for col in range(1, last_col + 1):
        ws.cell(3, col).fill = fill(BG)
        ws.cell(3, col).border = NONE


def section_label(ws, row, col, text, span=4):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    cell = ws.cell(row, col, text)
    cell.font = font(13, True, PRIMARY)
    cell.fill = fill(BG)
    cell.alignment = align("left", "center")
    cell.border = NONE
    for i in range(span):
        ws.cell(row, col + i).fill = fill(BG)
        ws.cell(row, col + i).border = NONE
    return cell


def header_row(ws, row, headers, fills=None):
    ws.row_dimensions[row].height = 22
    ws.auto_filter.ref = None  # set later
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row, i, h)
        cell.font = font(10, True, WHITE)
        cell.fill = fill(fills[i - 1] if fills else PRIMARY)
        cell.alignment = align("center", "center", True)
        cell.border = THIN
    return row


def style_data_row(ws, row, ncols, alt=False, height=18):
    ws.row_dimensions[row].height = height
    bg = ROW_ALT if alt else WHITE
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = fill(bg)
        cell.font = font(10, False, TEXT)
        cell.alignment = align("left", "center")
        cell.border = THIN


def kpi_card(ws, r, c, label, value, sub=None, accent=ACCENT, span=2, rows=3):
    """Paint a 3-row KPI card starting at (r,c). value may be a formula string."""
    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
    ws.merge_cells(start_row=r + 1, start_column=c, end_row=r + 1, end_column=c + span - 1)
    if rows >= 3:
        ws.merge_cells(start_row=r + 2, start_column=c, end_row=r + 2, end_column=c + span - 1)
    top = ws.cell(r, c, label)
    top.font = font(9, True, accent)
    top.fill = fill(WHITE)
    top.alignment = align("left", "center")
    mid = ws.cell(r + 1, c, value)
    mid.font = font(20, True, PRIMARY)
    mid.fill = fill(WHITE)
    mid.alignment = align("left", "center")
    if rows >= 3:
        bot = ws.cell(r + 2, c, sub or "")
        bot.font = font(8, False, TEXT_MUTED)
        bot.fill = fill(WHITE)
        bot.alignment = align("left", "center")
    for rr in range(r, r + rows):
        for cc in range(c, c + span):
            cell = ws.cell(rr, cc)
            cell.fill = fill(WHITE)
            cell.border = Border(
                left=Side(style="thin", color=SECONDARY_DEEP),
                right=Side(style="thin", color=SECONDARY_DEEP),
                top=Side(style="thin", color=SECONDARY if rr > r else accent),
                bottom=Side(style="thin", color=SECONDARY_DEEP),
            )
        ws.cell(rr, c).border = Border(
            left=Side(style="medium", color=accent),
            right=Side(style="thin", color=SECONDARY_DEEP),
            top=Side(style="thin", color=SECONDARY_DEEP),
            bottom=Side(style="thin", color=SECONDARY_DEEP),
        )
    ws.row_dimensions[r].height = 16
    ws.row_dimensions[r + 1].height = 28
    if rows >= 3:
        ws.row_dimensions[r + 2].height = 16
    return mid


def input_note(ws, row, col, text, span=8):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    cell = ws.cell(row, col, text)
    cell.font = font(9, False, TEXT_MUTED, italic=True)
    cell.fill = fill(BG)
    cell.alignment = align("left", "center", True)
    cell.border = NONE
    return cell


def freeze(ws, cell="A5"):
    ws.freeze_panes = cell


def auto_filter(ws, row, ncols, last_data_row):
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{last_data_row}"


def ivory_sheet(ws, max_row=80, max_col=16):
    paint_range(ws, 1, 1, max_row, max_col, BG)
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c).border = NONE


STATUS_COLORS = {
    "Active": SUCCESS,
    "Planning": INFO,
    "On Hold": WARNING,
    "Completed": PRIMARY,
    "Cancelled": DANGER,
    "To Do": TEXT_MUTED,
    "Scheduled": WARNING,
    "In Progress": ACCENT,
    "Inspection": INFO,
    "Blocked": DANGER,
    "Required": TEXT_MUTED,
    "Quoted": INFO,
    "Approved": SUCCESS,
    "Ordered": WARNING,
    "Shipped": ACCENT,
    "Delivered": INFO,
    "Installed": SUCCESS,
    "Paid": SUCCESS,
    "Committed": WARNING,
    "Planned": TEXT_MUTED,
    "Overdue": DANGER,
    "Open": ACCENT,
    "New": INFO,
    "Accepted": SUCCESS,
    "Dismissed": TEXT_MUTED,
    "Requested": TEXT_MUTED,
    "Received": INFO,
    "Compared": WARNING,
    "Rejected": DANGER,
    "Pass": SUCCESS,
    "Fail": DANGER,
    "Conditional": WARNING,
    "Pending": WARNING,
    "Signed": SUCCESS,
    "Expired": DANGER,
}


def status_fill(status: str) -> str:
    return STATUS_COLORS.get(status, PRIMARY_SOFT)
